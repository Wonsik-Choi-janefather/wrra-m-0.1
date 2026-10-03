from pathlib import Path
import json,copy,subprocess,re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
exec((ROOT/'content.py').read_text())
pages=json.loads((ROOT/'paper_content.json').read_text())
exec((ROOT/'review_edits.py').read_text())
for p in pages:
 for b in p['blocks']:
  if b['kind']=='p':
   b['en']=b['en'].replace('for one million addresses','at address cutoff N=10⁶')
   if b['en'].startswith('Joint calibration gives C=13.'):
    b['en']='Here b=μNB is field energy, τp=+1 and τn=−1; DB is the projected single-slot magnetic response and c0 its calibrated isoscalar correction. The quantity e− is the lowest energy of the two-state Δ matrix before the mass offset. '+b['en']
    b['ko']='b=μNB는 자기장 에너지이고 τp=+1, τn=−1이다. DB는 투영한 단일 슬롯 자기 응답이며 c0는 보정 등스칼라 항이다. e−는 질량 오프셋 전 두 상태 Δ 행렬의 최저 에너지다. '+b['ko']
   b['en']=b['en'].replace('A prime-family log-intensity εp','Writing νp(n) for the exponent of prime p in n, a prime-family log-intensity εp')
   b['ko']=b['ko'].replace('소수 가족 로그 강도 εp는','νp(n)는 n의 소인수 p의 지수다. 소수 가족 로그 강도 εp는')
   if b['en'].startswith('The three-body decay uses exact'):
    b['en']='The inherited inputs are GF=1.1663787×10⁻⁵ GeV⁻², |Vud|=0.97367 and Coulomb-weighted allowed phase-space integral I0=0.0589405883156 MeV⁵; GF is converted to MeV⁻² in the rate. '+b['en']
    b['ko']='계승 입력은 GF=1.1663787×10⁻⁵ GeV⁻², |Vud|=0.97367과 Coulomb 가중 허용 위상 적분 I0=0.0589405883156 MeV⁵다. 붕괴율에서 GF를 MeV⁻²로 변환한다. '+b['ko']
  elif b['kind']=='eq':
   b['tex']=b['tex'].replace(r'\sum_{3^k\le N}',r'\sum_{k\ge1,\,3^k\le N}').replace(r'\sum_{5^k\le N}',r'\sum_{k\ge1,\,5^k\le N}')
OUT=Path('output/upstream_integrated_r1');OUT.mkdir(parents=True,exist_ok=True)
eqs=[b for p in pages for b in p['blocks'] if b['kind']=='eq']
for b in eqs:b['tex']=re.sub(r'\{\\rm ([^{}]+)\}',lambda m:r'\mathrm{'+m.group(1)+'}',b['tex'])
(ROOT/'equations.md').write_text('\n\n'.join('$$\n'+b['tex']+'\n$$' for b in eqs))
subprocess.run(['pandoc',str(ROOT/'equations.md'),'-o',str(ROOT/'equations.docx')],check=True)
with ZipFile(ROOT/'equations.docx') as z:tree=etree.fromstring(z.read('word/document.xml'))
maths=tree.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})
assert len(maths)==len(eqs)
def font(r,sz=10.5,bold=False):
 r.font.name='NanumGothic';r.font.size=Pt(sz);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
 r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
def make_table(d,b,lang):
 d.add_paragraph(b[lang],'Caption')
 n=len(b['heads']);t=d.add_table(rows=1,cols=n);t.autofit=False
 widths=([3.25,3.45] if n==2 else [1.5,2.5,2.7] if n==3 else [3.25,1.15,1.15,1.15])
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for i,vals in enumerate([b['heads']]+b['rows']):
  row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,val in enumerate(vals):
   if ' / ' in val and lang=='en' and any('\uac00'<=c<='\ud7a3' for c in val):val=val.split(' / ')[0]
   elif ' / ' in val and lang=='ko' and any('\uac00'<=c<='\ud7a3' for c in val):val=val.split(' / ')[-1]
   c=row.cells[j];c.width=Inches(widths[j]);pr=c._tc.get_or_add_tcPr()
   if i==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EFF5');pr.append(sh)
   p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.1
   font(p.add_run(val),9,i==0)
 d.add_paragraph().paragraph_format.space_after=Pt(0)
for lang in ['ko','en']:
 d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
 s.left_margin=s.right_margin=Inches(.78);s.top_margin=Inches(.7);s.bottom_margin=Inches(.68)
 for name in ['Normal','Title','Subtitle','Heading 1','Caption']:
  st=d.styles[name];st.font.name='NanumGothic';st.font.color.rgb=RGBColor(0,0,0)
  st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
  st.font.italic=False
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=Pt(15.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Title'].font.size=Pt(19);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(7)
 d.styles['Subtitle'].font.size=Pt(10.5)
 d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 1'].font.bold=True
 d.styles['Heading 1'].paragraph_format.space_before=Pt(8);d.styles['Heading 1'].paragraph_format.space_after=Pt(9)
 d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].paragraph_format.keep_with_next=True
 title={'en':'Finite Upstream Generation and Common Energy–Current Readout in WRRA M','ko':'WRRA M의 유한 상류 생성과 공통 에너지·전류 판독'}[lang]
 d.add_paragraph(title,'Title')
 d.add_paragraph('Wonsik Choi · Jeongin Choi' if lang=='en' else '최원식 Wonsik Choi · 최정인 Jeongin Choi')
 d.add_paragraph('Wonsik Choi: Independent researcher, Seoul, Republic of Korea' if lang=='en' else '최원식: 독립 연구자, 대한민국 서울','Caption')
 d.add_paragraph('Correspondence: janefather@gmail.com · ORCID 0009-0001-4263-9772','Caption')
 d.add_paragraph('Integrated manuscript 1.0-r1 · Upstream frozen at 0.10 · 3 October 2026' if lang=='en' else '통합 원고 1.0-r1 · 상류 0.10 마무리 · 2026년 10월 3일','Subtitle')
 d.add_paragraph('DOI: 10.5281/zenodo.23119041','Caption')
 mi=0;md=[]
 for idx,p in enumerate(pages):
  h=d.add_paragraph(p[lang],'Heading 1');h.paragraph_format.page_break_before=idx>0
  md.append('# '+p[lang])
  for b in p['blocks']:
   if b['kind']=='p':
    txt=b[lang]
    if idx==0 and not md[-1].startswith('Keywords') and 'We present' in txt:txt=txt.replace('WRRA M upstream sequence','Wonsik Reality Renderer Architecture (WRRA) M upstream sequence')
    if txt.startswith(('Statements and Declarations.','선언.')):
     txt=txt.split(' Current-manuscript')[0] if lang=='en' else txt.split(' 이번 원고')[0]
    pp=d.add_paragraph(txt);md.append(txt)
    if idx==11:pp.paragraph_format.line_spacing=Pt(13);pp.paragraph_format.space_after=Pt(6)
   elif b['kind']=='table':make_table(d,b,lang)
   else:
    node=copy.deepcopy(maths[mi]);mi+=1
    for dp in node.xpath('.//*[local-name()="dPr"]'):
     order={'begChr':0,'sepChr':1,'endChr':2,'grow':3,'shp':4,'ctrlPr':5}
     children=sorted(list(dp),key=lambda x:order.get(etree.QName(x).localname,6))
     for x in list(dp):dp.remove(x)
     for x in children:dp.append(x)
    for mr in node.xpath('.//*[local-name()="r"]'):
     pr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Latin Modern Math');rf.set(qn('w:hAnsi'),'Latin Modern Math');pr.append(rf)
     sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);mr.insert(1 if len(mr) and etree.QName(mr[0]).localname=='rPr' else 0,pr)
    pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=1;pp.paragraph_format.keep_together=True
    pp.paragraph_format.space_after=Pt(8);pp._p.append(node);font(pp.add_run('  ('+str(mi)+')'),9)
    md.append('$$\n'+b['tex']+'\n$$')
 fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=fp.add_run();fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'begin');r._r.append(fld)
 instr=OxmlElement('w:instrText');instr.text=' PAGE ';r._r.append(instr)
 fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'separate');r._r.append(fld)
 font(fp.add_run('1'),8)
 fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'end');fp.add_run()._r.append(fld)
 for tree in [d.styles._element,d._element]:
  for el in tree.xpath('.//w:pBdr'):el.getparent().remove(el)
 d.core_properties.title=title;d.core_properties.author='Wonsik Choi; Jeongin Choi';d.core_properties.subject='Integrated upstream 0.10 research manuscript'
 stem='WRRA_M_Upstream_Integrated_1_0_r1_'+lang.upper()+'_2026_10_03'
 d.save(OUT/(stem+'.docx'));(ROOT/(stem+'.md')).write_text('\n\n'.join(md))
 print(OUT/(stem+'.docx'))
print('Native equations:',len(eqs))

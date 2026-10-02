from pathlib import Path
import copy
import json
import subprocess
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from note import TITLE,SUBTITLE,B,R

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'paper';OUT.mkdir(exist_ok=True)
ASSETS=ROOT/'assets';ASSETS.mkdir(exist_ok=True)
STEM='WRRA_M_Spatial_Stability_Excitation_Scale_v0_5_KO_2026_10_02'

def font(run,size=None,bold=False):
 run.font.name='NanumGothic';run.font.bold=bold
 run.font.color.rgb=RGBColor(0,0,0)
 run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
 if size:run.font.size=Pt(size)

def plots():
 fig,axes=plt.subplots(1,2,figsize=(9,3.1))
 bare=R['bare_instability']['shell_sequence'];frozen=R['frozen_parameter_convergence']
 axes[0].plot([x['shell_order'] for x in bare],[x['ground_energy_over_delta'] for x in bare],'o-',color='#a33b30')
 axes[0].set(xlabel='Shell order K',ylabel='Ground energy / delta',title='Unbounded quadratic link')
 axes[1].plot([x['shell_order'] for x in frozen],[x['gap_MeV'] for x in frozen],'o-',color='#235d91')
 axes[1].axhline(R['reference']['gap_MeV'],color='#555555',ls='--',lw=1)
 axes[1].set(xlabel='Shell order K',ylabel='Gap (MeV)',title='Bounded link, fixed parameters')
 for ax in axes:ax.grid(alpha=.2);ax.set_xticks([1,2,4,6,8])
 fig.tight_layout();fig.savefig(ASSETS/'convergence.png',dpi=180,bbox_inches='tight');plt.close(fig)

def add_table(doc,b):
 cap=doc.add_paragraph(b['caption'],'Caption');cap.paragraph_format.keep_with_next=True
 tab=doc.add_table(rows=1,cols=len(b['headers']));tab.autofit=False
 borders=OxmlElement('w:tblBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  v=OxmlElement('w:'+edge);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
 tab._tbl.tblPr.append(borders)
 repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
 for j,width in enumerate(b['widths']):tab.columns[j].width=Inches(width)
 for i,values in enumerate([b['headers']]+b['rows']):
  row=tab.rows[0] if i==0 else tab.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,value in enumerate(values):
   cell=row.cells[j];cell.width=Inches(b['widths'][j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   props=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ('top','bottom','left','right'):
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'60' if side in ('top','bottom') else '100');e.set(qn('w:type'),'dxa');m.append(e)
   props.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');props.append(sh)
   pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.08
   pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   if len(str(value))>24:pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
   font(pp.add_run(str(value)),9.5,i==0)
 pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(4);font(pp.add_run(' '),3)

def build():
 plots()
 (ROOT/'note.json').write_text(json.dumps({'title':TITLE,'subtitle':SUBTITLE,'blocks':B},ensure_ascii=False,indent=2),encoding='utf8')
 eqs=[b for b in B if b['kind']=='eq']
 math_in=ASSETS/'equations.txt'
 math_in.write_text('\n\n'.join('$$\n'+b['tex']+'\n$$' for b in eqs),encoding='utf8')
 subprocess.run(['pandoc','-f','markdown',str(math_in),'-o',str(ASSETS/'equations.docx')],check=True)
 with ZipFile(ASSETS/'equations.docx') as z:tree=etree.fromstring(z.read('word/document.xml'))
 native=tree.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})
 assert len(native)==len(eqs)
 maths={b['label']:copy.deepcopy(el) for b,el in zip(eqs,native)}
 d=Document();s=d.sections[0]
 s.page_width=Inches(8.5);s.page_height=Inches(11)
 s.left_margin=s.right_margin=Inches(.8);s.top_margin=Inches(.7);s.bottom_margin=Inches(.65)
 s.header_distance=s.footer_distance=Inches(.25)
 for name in ('Normal','Title','Subtitle','Heading 1','Caption'):
  st=d.styles[name];st.font.name='NanumGothic';st.font.color.rgb=RGBColor(0,0,0)
  st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
  st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;st.paragraph_format.first_line_indent=Pt(0)
  st.font.italic=False
  for e in st._element.xpath('.//w:pBdr|.//w:numPr'):e.getparent().remove(e)
 n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=Pt(16.5);n.paragraph_format.space_after=Pt(6)
 d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(8)
 d.styles['Subtitle'].font.size=Pt(11);d.styles['Subtitle'].paragraph_format.space_after=Pt(8)
 st=d.styles['Heading 1'];st.font.size=Pt(13);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
 st=d.styles['Caption'];st.font.size=Pt(9);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.05
 hp=s.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(hp.add_run('WRRA M 내부 공간 안정성·여기 척도 0.5   최원식'),8)
 fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=fp.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
 inst=OxmlElement('w:instrText');inst.text=' PAGE ';r._r.append(inst)
 a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a);font(fp.add_run('1'),8)
 a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');fp.add_run()._r.append(a)
 d.add_paragraph(TITLE,'Title');d.add_paragraph(SUBTITLE,'Subtitle')
 pp=d.add_paragraph('최원식 Wonsik Choi   Seoul   2026년 10월 2일');pp.paragraph_format.space_after=Pt(2)
 pp=d.add_paragraph('ORCID 0009-0001-4263-9772   janefather@gmail.com');pp.paragraph_format.space_after=Pt(6)
 next_page=False
 for b in B:
  kind=b['kind']
  if kind=='p':d.add_paragraph(b['text'])
  elif kind=='h':
   hp=d.add_paragraph(b['text'],'Heading 1')
   if next_page:hp.paragraph_format.page_break_before=True;next_page=False
  elif kind=='break':next_page=True
  elif kind=='table':add_table(d,b)
  elif kind=='fig':
   pp=d.add_paragraph();pp.paragraph_format.line_spacing=1;pp.paragraph_format.keep_with_next=True
   pp.add_run().add_picture(str(ASSETS/(b['name']+'.png')),width=Inches(b['width']))
   d.add_paragraph(b['caption'],'Caption')
  elif kind=='eq':
   pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=1
   pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.space_after=Pt(7);pp.paragraph_format.keep_together=True
   node=maths[b['label']]
   for mr in node.xpath('.//*[local-name()="r"]'):
    props=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Latin Modern Math');fonts.set(qn('w:hAnsi'),'Latin Modern Math');props.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');props.append(sz);mr.insert(0,props)
   pp._p.append(node);font(pp.add_run('  ('+b['label']+')'),9)
 d.core_properties.title=TITLE;d.core_properties.subject=SUBTITLE;d.core_properties.author='Wonsik Choi'
 path=OUT/(STEM+'.docx');d.save(path)
 md=['# '+TITLE,'',SUBTITLE,'','최원식 Wonsik Choi · 2026-10-02','']
 for b in B:
  if b['kind']=='h':md+=['## '+b['text'],'']
  elif b['kind']=='p':md += [b['text'],'']
  elif b['kind']=='eq':md+=['$$',b['tex'],'$$','('+b['label']+')','']
  elif b['kind']=='table':
   md += [b['caption'],'','| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
   md += ['| '+' | '.join(row)+' |' for row in b['rows']];md+=['']
  elif b['kind']=='fig':md+=['!['+b['caption']+'](source/assets/'+b['name']+'.png)','']
 (ROOT.parent/'manuscript_KO.md').write_text('\n'.join(md),encoding='utf8')
 print(path)

if __name__=='__main__':build()

from pathlib import Path
import json
import os
import subprocess
import sys
import copy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from matplotlib.font_manager import FontProperties
from manuscript import TITLE, SUBTITLE, BLOCKS

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
ASSETS=ROOT/'assets'; ASSETS.mkdir(exist_ok=True)
STEM='WRRA_M_Upper_Two_Stage_Filter_Hypothesis_v1_0_KO_2026_10_01'
def architecture():
    font=FontProperties(fname='/usr/local/share/fonts/wrra/NotoSansCJKkr-Regular.otf')
    fig,ax=plt.subplots(figsize=(8.5,4.1),dpi=250)
    ax.set_xlim(0,10);ax.set_ylim(0,5);ax.axis('off')
    boxes={
      'source':(0.3,3.75,2.0,.75,'SOURCE\n방출 가능성'),
      'in':(3.0,3.75,2.7,.75,'1단계 진입 필터\n초기 변동'),
      'return':(7.1,3.1,2.4,.75,'상류 복귀 R\n반사 및 복귀'),
      'zero':(3.0,2.2,2.7,.75,'2단계 0차원 복귀 필터\n정상화 후 접힘 판정'),
      'actual':(3.0,.95,2.7,.75,'Actual 잔존\n접힘 보존'),
      'read':(7.1,.6,2.4,1.1,'잔존의 판독\n표현형 φ  비표현형 D')}
    for key,(x,y,w,h,t) in boxes.items():
      ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.05',facecolor='white',edgecolor='black',linewidth=1.1))
      ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=10,fontproperties=font)
    def arrow(a,b,label=None,lp=None):
      ax.annotate('',xy=b,xytext=a,arrowprops=dict(arrowstyle='->',lw=1.2,color='black'))
      if label:ax.text(*lp,label,fontproperties=font,fontsize=9,ha='center',va='center',bbox=dict(facecolor='white',edgecolor='none',pad=1))
    arrow((2.34,4.125),(2.94,4.125))
    arrow((4.35,3.71),(4.35,3.0),'진입',(4.73,3.34))
    arrow((5.74,4.1),(7.05,3.62),'반사',(6.55,4.2))
    arrow((5.74,2.58),(7.05,3.25),'무접힘',(6.55,2.69))
    arrow((4.35,2.16),(4.35,1.75),'복귀 차단',(5.0,1.95))
    arrow((5.74,1.32),(7.05,1.17),'판독',(6.4,1.52))
    ax.text(.42,.27,'두 필터는 통과와 복귀를 판정한다   판독은 잔존 상태의 출력이다',fontproperties=font,fontsize=9)
    fig.subplots_adjust(left=.015,right=.995,bottom=.025,top=.98)
    fig.savefig(ASSETS/'architecture.png',bbox_inches='tight',pad_inches=.06);plt.close(fig)

def set_font(run,size=None,bold=None):
    run.font.name='Noto Sans CJK KR'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
    run.font.color.rgb=RGBColor(0,0,0)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold

def no_split(row):
    trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))

def table(doc,b):
    caption=doc.add_paragraph(b['caption'],'Caption');caption.paragraph_format.keep_with_next=True
    tab=doc.add_table(rows=1,cols=len(b['headers']));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    for i,w in enumerate(b['widths']):tab.columns[i].width=Inches(w*.99)
    pr=tab._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
       x=OxmlElement('w:'+name);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders)
    repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
    data=[b['headers']]+b['rows']
    for row_idx,values in enumerate(data):
      row=tab.rows[0] if row_idx==0 else tab.add_row();no_split(row)
      for j,value in enumerate(values):
       cell=row.cells[j];cell.width=Inches(b['widths'][j]*.99);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
       cpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
       for name in ('top','bottom','left','right'):
        el=OxmlElement('w:'+name);el.set(qn('w:w'),'85' if name in ('top','bottom') else '100');el.set(qn('w:type'),'dxa');m.append(el)
       cpr.append(m)
       if row_idx==0:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');cpr.append(sh)
       pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=1.12
       pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(str(value))>25 else WD_ALIGN_PARAGRAPH.CENTER
       rr=pp.add_run(str(value));set_font(rr,9.5,row_idx==0)
    after=doc.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.space_before=Pt(0);after.paragraph_format.line_spacing=1
    after.add_run().font.size=Pt(3)

def build():
    (ROOT/'manuscript.json').write_text(json.dumps({'title':TITLE,'subtitle':SUBTITLE,'author':'Wonsik Choi','date':'2026-10-01','blocks':BLOCKS},ensure_ascii=False,indent=2),encoding='utf8')
    equations=[b for b in BLOCKS if b['kind']=='eq']
    math_input=ASSETS/'math_input.txt'
    math_input.write_text('\n\n'.join('$$\n'+b['tex']+'\n$$' for b in equations),encoding='utf8')
    math_doc=ASSETS/'native_math.docx'
    subprocess.run(['pandoc','-f','markdown',str(math_input),'-o',str(math_doc)],check=True)
    with ZipFile(math_doc) as z:
      tree=etree.fromstring(z.read('word/document.xml'))
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    native=tree.xpath('//m:oMath',namespaces=ns)
    assert len(native)==len(equations),(len(native),len(equations))
    math_nodes={b['label']:copy.deepcopy(node) for b,node in zip(equations,native)}
    architecture()
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(.8);s.top_margin=Inches(.75);s.bottom_margin=Inches(.7)
    s.header_distance=s.footer_distance=Inches(.3)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
     st=d.styles[name];st.font.name='Noto Sans CJK KR';st.font.color.rgb=RGBColor(0,0,0)
     st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
     st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
     st.font.italic=False
     for el in st._element.xpath('.//w:pBdr|.//w:numPr'):
       el.getparent().remove(el)
    n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=Pt(16.5);n.paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(21);d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(9)
    d.styles['Subtitle'].font.size=Pt(11.5);d.styles['Subtitle'].paragraph_format.space_after=Pt(11)
    for name,size in (('Heading 1',15),('Heading 2',11.5)):
     st=d.styles[name];st.font.size=Pt(size);st.font.bold=True
     st.paragraph_format.space_before=Pt(14 if name=='Heading 1' else 10);st.paragraph_format.space_after=Pt(6)
     st.paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False;d.styles['Caption'].paragraph_format.space_after=Pt(6)
    header=s.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run('WRRA M 상류 가설 1.0   최원식'),8)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' PAGE ';run._r.append(instr)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');run._r.append(sep)
    cached=footer.add_run('1');set_font(cached,8)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');footer.add_run()._r.append(end)
    d.add_paragraph(TITLE,'Title');d.add_paragraph(SUBTITLE,'Subtitle')
    for line in ('최원식 Wonsik Choi','Independent Researcher  Seoul  Republic of Korea','ORCID 0009-0001-4263-9772   janefather@gmail.com','상류 구조 가설 1.0   2026년 10월 1일'):
      pp=d.add_paragraph();pp.paragraph_format.space_after=Pt(2);set_font(pp.add_run(line),9)
    for b in BLOCKS:
     if b['kind']=='h':d.add_paragraph(b['text'],'Heading '+str(b['level']))
     elif b['kind']=='p':d.add_paragraph(b['text'])
     elif b['kind']=='table':table(d,b)
     elif b['kind']=='eq':
      pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(8)
      pp.paragraph_format.line_spacing=1
      pp.paragraph_format.keep_together=True
      node=math_nodes[b['label']]
      for mr in node.xpath('.//*[local-name()="r"]'):
        rpr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Latin Modern Math');fonts.set(qn('w:hAnsi'),'Latin Modern Math');rpr.append(fonts)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');rpr.append(sz);mr.insert(0,rpr)
      pp._p.append(node)
      set_font(pp.add_run('   ('+b['label']+')'),9)
     elif b['kind']=='fig':
      pp=d.add_paragraph();pp.paragraph_format.keep_with_next=True
      pp.paragraph_format.line_spacing=1;pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
      pp.add_run().add_picture(str(ASSETS/(b['name']+'.png')),width=Inches(6.1))
      d.add_paragraph(b['caption'],'Caption')
    d.core_properties.title=TITLE;d.core_properties.subject=SUBTITLE;d.core_properties.author='Wonsik Choi';d.core_properties.keywords='WRRA; two-stage dimensional filter; residue; frame update'
    path=OUT/(STEM+'.docx');d.save(path);print(path)

if __name__=='__main__':build()

from pathlib import Path
import copy
import json
import subprocess
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from note import TITLE,SUBTITLE,B,R

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent/'paper';OUT.mkdir(exist_ok=True)
ASSETS=ROOT/'assets';ASSETS.mkdir(exist_ok=True)
STEM='WRRA_M_Internal_Spatial_Coupling_Joint_Currents_v0_4_r1_KO_2026_10_03'

def font(run,size=None,bold=False):
 run.font.name='NanumGothic';run.font.bold=bold
 run.font.color.rgb=RGBColor(0,0,0)
 run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
 if size:run.font.size=Pt(size)

def plots():
 pass

def add_table(doc,b):
 cap=doc.add_paragraph(b['caption'],'Caption');cap.paragraph_format.keep_with_next=True
 tab=doc.add_table(rows=1,cols=len(b['headers']));tab.autofit=False
 borders=OxmlElement('w:tblBorders')
 for edge in ('top','left','bottom','right','insideH','insideV'):
  v=OxmlElement('w:'+edge);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
 tab._tbl.tblPr.append(borders)
 repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
 for j,width in enumerate(b['widths']):tab.columns[j].width=Inches(width)
 for i,values in enumerate([b['headers']]+b['rows']):
  row=tab.rows[0] if i==0 else tab.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for j,value in enumerate(values):
   cell=row.cells[j];cell.width=Inches(b['widths'][j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   props=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side in ('top','bottom','left','right'):
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'60' if side in ('top','bottom') else '100');e.set(qn('w:type'),'dxa');m.append(e)
   props.append(m)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');props.append(sh)
   pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.08
   pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   if len(str(value))>24:pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
   font(pp.add_run(str(value)),9.5,i==0)
 pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(4);font(pp.add_run(' '),3)

def build():
 plots()
 (ROOT/'note.json').write_text(json.dumps({'title':TITLE,'subtitle':SUBTITLE,'blocks':B},ensure_ascii=False,indent=2),encoding='utf8')
 eqs=[b for b in B if b['kind']=='eq']
 math_in=ASSETS/'equations.txt'
 math_in.write_text('\n\n'.join('$$\n'+b['tex']+'\n$$' for b in eqs),encoding='utf8')
 subprocess.run(['pandoc','-f','markdown',str(math_in),'-o',str(ASSETS/'equations.docx')],check=True)
 with ZipFile(ASSETS/'equations.docx') as z:tree=etree.fromstring(z.read('word/document.xml'))
 native=tree.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})
 assert len(native)==len(eqs)
 maths={b['label']:copy.deepcopy(el) for b,el in zip(eqs,native)}
 d=Document();s=d.sections[0]
 s.page_width=Inches(8.5);s.page_height=Inches(11)
 s.left_margin=s.right_margin=Inches(.8);s.top_margin=Inches(.7);s.bottom_margin=Inches(.65)
 s.header_distance=s.footer_distance=Inches(.25)
 for name in ('Normal','Title','Subtitle','Heading 1','Caption'):
  st=d.styles[name];st.font.name='NanumGothic';st.font.color.rgb=RGBColor(0,0,0)
  st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumGothic')
  st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;st.paragraph_format.first_line_indent=Pt(0)
  st.font.italic=False
  for e in st._element.xpath('.//w:pBdr|.//w:numPr'):e.getparent().remove(e)
 n=d.styles['Normal'];n.font.size=Pt(11);n.paragraph_format.line_spacing=Pt(16.5);n.paragraph_format.space_after=Pt(6)
 d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True;d.styles['Title'].paragraph_format.space_after=Pt(8)
 d.styles['Subtitle'].font.size=Pt(11);d.styles['Subtitle'].paragraph_format.space_after=Pt(8)
 st=d.styles['Heading 1'];st.font.size=Pt(13);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
 st=d.styles['Caption'];st.font.size=Pt(9);st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.05
 hp=s.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(hp.add_run('WRRA M 내부 공간 결합과 공동 전류 0.4   최원식'),8)
 fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 r=fp.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');r._r.append(a)
 inst=OxmlElement('w:instrText');inst.text=' PAGE ';r._r.append(inst)
 a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'separate');r._r.append(a);font(fp.add_run('1'),8)
 a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'end');fp.add_run()._r.append(a)
 d.add_paragraph(TITLE,'Title');d.add_paragraph(SUBTITLE,'Subtitle')
 pp=d.add_paragraph('최원식 Wonsik Choi   Seoul   2026년 10월 3일');pp.paragraph_format.space_after=Pt(2)
 pp=d.add_paragraph('ORCID 0009-0001-4263-9772   janefather@gmail.com');pp.paragraph_format.space_after=Pt(6)
 d.add_paragraph('교정 자료집 DOI 10.5281/zenodo.23112253', 'Caption')
 next_page=False
 for b in B:
  kind=b['kind']
  if kind=='p':d.add_paragraph(b['text'])
  elif kind=='h':
   hp=d.add_paragraph(b['text'],'Heading 1')
   if next_page:hp.paragraph_format.page_break_before=True;next_page=False
  elif kind=='break':next_page=True
  elif kind=='table':add_table(d,b)
  elif kind=='fig':
   pp=d.add_paragraph();pp.paragraph_format.line_spacing=1;pp.paragraph_format.keep_with_next=True
   pp.add_run().add_picture(str(ASSETS/(b['name']+'.png')),width=Inches(b['width']))
   d.add_paragraph(b['caption'],'Caption')
  elif kind=='eq':
   pp=d.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=1
   pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.space_after=Pt(7);pp.paragraph_format.keep_together=True
   node=maths[b['label']]
   for dp in node.xpath('.//*[local-name()="dPr"]'):
    order={'begChr':0,'sepChr':1,'endChr':2,'grow':3,'shp':4,'ctrlPr':5}
    children=sorted(list(dp),key=lambda x:order.get(etree.QName(x).localname,6))
    for child in list(dp):dp.remove(child)
    for child in children:dp.append(child)
   for mr in node.xpath('.//*[local-name()="r"]'):
    props=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Latin Modern Math');fonts.set(qn('w:hAnsi'),'Latin Modern Math');props.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');props.append(sz);mr.insert(1 if len(mr) and etree.QName(mr[0]).localname=='rPr' else 0,props)
   pp._p.append(node);font(pp.add_run('  ('+b['label']+')'),9)
 d.core_properties.title=TITLE;d.core_properties.subject=SUBTITLE;d.core_properties.author='Wonsik Choi'
 path=OUT/(STEM+'.docx');d.save(path);print(path)

if __name__=='__main__':build()

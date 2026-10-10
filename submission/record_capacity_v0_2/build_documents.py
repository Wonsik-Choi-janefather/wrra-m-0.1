"""Build bilingual DOCX files. Install Noto Sans CJK KR for Korean rendering."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from matplotlib.mathtext import math_to_image
from matplotlib.font_manager import FontProperties

ROOT=Path(__file__).resolve().parent
EQUATIONS=[
 r'$|E|\geq\max_r |r^{-1}(r)|.$',
 r'$f_i=1+\frac{\ln n_i}{4\ln N},\qquad w_i=\frac{n_i^{-\alpha}}{\sum_j n_j^{-\alpha}}.$',
 r'$t=(f_2-f_3,f_3-f_1,f_1-f_2),\quad \delta_i=\frac{0.2(t_i/w_i)}{\max_j|t_j/w_j|}.$',
 r'$p_{i0}^{\pm}=w_i(0.6\pm\delta_i),\quad p_{i1}^{\pm}=w_i(0.4\mp\delta_i),\quad b_{ik}=2kf_i.$',
 r'$\Phi_{\pm}(x;J)=x^2-\ln\left[\sum_{ik}p_{ik}^{\pm}\exp(-xb_{ik})\right]-Jx.$',
 r'$\beta Q=\Delta S_S+D(u\Vert\tau).$'
]
def set_font(style,name,size):
    style.font.name=name; style.font.size=Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
def rich(p,text):
    for j,part in enumerate(text.split('**')):
        r=p.add_run(part);r.bold=bool(j%2)
def build(lang):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.1);s.bottom_margin=Cm(2.1);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    normal=d.styles['Normal'];set_font(normal,'DejaVu Serif' if lang=='en' else 'Noto Sans CJK KR',10.5)
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.line_spacing=1.15 if lang=='en' else 1.25
    normal.paragraph_format.widow_control=True
    # Remove decorative title borders inherited from the runtime template.
    for node in list(d.styles.element.xpath('.//w:pBdr')):
        node.getparent().remove(node)
    for name,size in [('Title',20),('Subtitle',12),('Heading 1',12),('Heading 2',11),('Caption',9)]:
        st=d.styles[name];set_font(st,'DejaVu Sans' if lang=='en' else 'Noto Sans CJK KR',size)
        st.font.color.rgb=RGBColor.from_string('1E2933')
        st.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'):
            st.paragraph_format.space_before=Pt(14);st.paragraph_format.keep_with_next=True
    d.core_properties.title='Record Capacity and Irreversibility in a Finite-State Model' if lang=='en' else '유한 상태 모형에서 기록 용량과 비가역성'
    d.core_properties.author='Wonsik Choi; Jeongin Choi'
    d.core_properties.subject='WRRA finite record capacity and reset, v0.2'
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
    eqnum=0
    source=(ROOT/f'manuscript_{lang}.md').read_text()
    source=re.sub(r'(?m)^(#{1,3} [^\n]+)\n',r'\n\n\1\n\n',source)
    source=source.replace('\nEQ:','\n\nEQ:')
    for block in re.split(r'\n\s*\n',source):
        block=block.strip()
        if not block:continue
        if block.startswith('### '):d.add_paragraph(block[4:],'Heading 1')
        elif block.startswith('## '):d.add_paragraph(block[3:],'Subtitle')
        elif block.startswith('# '):d.add_paragraph(block[2:],'Title')
        elif block in ('FIGURE','FIGURE_ACCUMULATION'):
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(ROOT/('figure.png' if block=='FIGURE' else 'accumulation.png')),width=Cm(16.5))
            shape._inline.docPr.set('descr','Exact history versus parity residue state counts and finite thermal reset ledger')
        elif block.startswith('Caption: '):
            p=d.add_paragraph(block[9:],'Caption')
        elif block.startswith('EQ: '):
            eqnum+=1;fn=ROOT/f'equation_{eqnum}.png'
            math_to_image(EQUATIONS[eqnum-1],fn,prop=FontProperties(size=13),dpi=240,format='png',color='black')
            from PIL import Image
            with Image.open(fn) as im:width=min(5.85,im.width/240)
            table=d.add_table(rows=1,cols=2);table.autofit=False
            table.columns[0].width=Cm(15);table.columns[1].width=Cm(1.5)
            p=table.cell(0,0).paragraphs[0];p.paragraph_format.space_after=Pt(6)
            shape=p.add_run().add_picture(str(fn),width=Inches(width));shape._inline.docPr.set('descr',block[4:])
            p=table.cell(0,1).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run(f'({eqnum})')
            trpr=table.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        elif re.match(r'^\[\d+\] ',block):
            for line in block.splitlines():
                p=d.add_paragraph(line);p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.keep_with_next=not line.startswith('[6] ')
                for r in p.runs:r.font.size=Pt(9)
        else:
            p=d.add_paragraph();rich(p,block)
    d.save(ROOT/f'WRRA_Record_Capacity_v0.2_{lang.upper()}.docx')

if __name__=='__main__':
    for lang in ['en','ko']:build(lang)

from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
R=Path(__file__).resolve().parent.parent
p=R/'output/WRRA_Address_Particle_Interface_v0_2_EN.docx';d=Document(p)
for t in d.tables:
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for para in c.paragraphs:para.paragraph_format.keep_with_next=i<len(t.rows)-1
 d.save(p)

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
R=Path(__file__).resolve().parent.parent/'output'
for f in R.glob('*.docx'):
 d=Document(f)
 for root in [d.styles.element,d._element]:
  for e in list(root.iter(qn('w:pBdr'))):e.getparent().remove(e)
 for t in d.tables:
  for row in t.rows:
   for c in row.cells:
    for p in c.paragraphs:p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
 if '_KO_' in f.name:
  d.styles['Normal'].paragraph_format.line_spacing=Pt(16)
  d.styles['Title'].font.size=Pt(18)
  for s in ['Heading 1','Heading 2']:d.styles[s].paragraph_format.space_before=Pt(10);d.styles[s].paragraph_format.space_after=Pt(5)
 else:
  # Let the conclusion continue after the final scope paragraph instead of an orphan page.
  pp=d.paragraphs
  for i,p in enumerate(pp):
   if p.text.startswith('10 Assessment') and i:
    prev=pp[i-1]
    if prev._p.xpath('.//w:br'):prev._element.getparent().remove(prev._element)
  for mr in d._element.iter(qn('m:r')):
   rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');rp.append(sz);mr.insert(0,rp)
 d.save(f)

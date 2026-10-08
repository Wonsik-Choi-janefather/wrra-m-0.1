from pathlib import Path
import subprocess
from docx import Document
from copy import deepcopy
R=Path(__file__).resolve().parent;O=R.parent/'output'
latex=[r'\chi(n)=\frac{\sum_p\nu_p^2}{\Omega(n)^2},\qquad 0<\chi(n)\leq1.\quad(1)',r'|\psi_n\rangle=\sqrt{\chi(n)}|e\rangle+\sqrt{1-\chi(n)}|\mu\rangle.\quad(2)',r'V|n\rangle=|n\rangle_{\mathrm{record}}\otimes|\psi_n\rangle,\qquad V^\dagger V=I.\quad(3)',r'P_e=\sum_n\chi(n)|n\rangle\langle n|,\qquad P_\mu=I-P_e.\quad(4)',r'p_\mu c=0,\qquad p_e c=\sqrt{\varepsilon_\mu^2-\varepsilon_e^2}.\quad(5)',r'\langle E_{\mathrm{rest}}\rangle=2[X_d\varepsilon_e+(1-X_d)\varepsilon_\mu],\quad\langle E_{\mathrm{kin}}\rangle=E_*-\langle E_{\mathrm{rest}}\rangle.\quad(6)',r'E_a(v)=2\sqrt{\varepsilon_a^2+(p_a(1)c)^2v^{-2/3}}.\quad(7)',r'P_aV=\frac{2(p_a(1)c)^2v^{-2/3}}{3\sqrt{\varepsilon_a^2+(p_a(1)c)^2v^{-2/3}}}.\quad(8)',r'\mathcal E(v)=0.05[X_d\sqrt{r+(1-r)v^{-2/3}}+1-X_d]+0.268+0.682v.\quad(9)',r'\mathcal P(v)=pV/E_0=-v\mathcal E\prime(v),\qquad q_{\mathrm{inst}}=\frac12\left[1+\frac{3\mathcal P(1)}{\mathcal E(1)}\right].\quad(10)',r'q_{\mathrm{inst}}=-0.523+0.025X_d(1-r).\quad(11)']
(R/'math.txt').write_text('\n\n'.join('$$'+s+'$$' for s in latex))
subprocess.run(['pandoc',str(R/'math.txt'),'-f','markdown','-o',str(R/'math.docx')],check=True)
src=Document(R/'math.docx');eqs=src._element.xpath('.//m:oMathPara');d=Document(O/'WRRA_Address_Particle_Interface_v0_2_EN.docx');old=d._element.xpath('.//m:oMathPara');assert len(old)==len(eqs)==11
for a,b in zip(old,eqs):a.getparent().replace(a,deepcopy(b))
for p in d.paragraphs:
 if p.text.startswith('In (8),'):p.text=p.text.replace('pa,pressure','P_a')
d.save(O/'WRRA_Address_Particle_Interface_v0_2_EN.docx')

from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
R=Path(__file__).resolve().parent.parent
p=R/'output/WRRA_Address_Particle_Interface_v0_2_KO_Guide.docx';d=Document(p)
for para in list(d.paragraphs):
 if para._p.xpath('.//w:br[@w:type="page"]') and not para.text.strip():
  para._p.getparent().remove(para._p)
for t in d.tables:
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for para in c.paragraphs:para.paragraph_format.keep_with_next=i<len(t.rows)-1
for para in d.paragraphs:
 if para.text.startswith('첫째, 왜 이 소인수'):
  para.text='이 가능성을 실제 물리 모형으로 확장할 때에는 소인수 집중도와 입자 선택을 잇는 공통운반자의 연산자, 준비·측정 절차, 붕괴와 에너지 교환을 명세해야 한다. 이는 현재 토이모델의 계산 결과와 구분되는 후속 과제다. 지금은 15채널 전체를 이 코드에 통합했다고 주장하지 않는다.'
 if para.text.startswith('현재 문서는'):
  para.text=para.text.replace('실제 입자 채널의 판독량','모형의 입자 채널 판독량')
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   if '확률·상태·에너지 보존 실패' in c.text:c.text=c.text.replace('확률·상태·에너지 보존 실패','확률 정규화·출력 양자수·에너지 회계 실패')
for para in list(d.paragraphs):
 if para.text.startswith('영문 제목:'):
  para._p.getparent().remove(para._p)
d.save(p)

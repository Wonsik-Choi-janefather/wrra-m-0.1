from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parent.parent
pairs=[('Address Resolved Particle Readout in a Finite WRRA Model','From Number Structure to Particle Selection'),('Conservation conditions and the limits of aggregate energy matching','A Possible Readout and Its Energy–Pressure Constraints in a Finite WRRA Toy Model'),('산술 주소에서 입자 판독으로 이어지는 WRRA 미시 사상','수의 구조에서 입자 선택으로')]
for f in (R/'output').glob('*.docx'):
 d=Document(f)
 for p in d.paragraphs:
  for a,b in pairs:
   if a in p.text:
    for run in p.runs:run.text=run.text.replace(a,b)
 if 'KO_Guide' in f.name:
  d.paragraphs[1].text='유한한 WRRA 토이모델에서 가능한 입자 선택과 에너지·압력의 제약\n후속 논문 0.2 한글 해설 · 2026년 10월 8일'
 d.save(f)
f=R/'source/build.py';t=f.read_text()
for a,b in pairs:t=t.replace(a,b)
f.write_text(t)

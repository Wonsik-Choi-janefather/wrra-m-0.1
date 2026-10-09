from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

P=Path(__file__).resolve().parent
def build(source,stem):
    subprocess.run(['pandoc',str(P/source),'-o',str(P/(stem+'.docx')),'--resource-path',str(P)],check=True)
    d=Document(P/(stem+'.docx'))
    sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
    for name in ['Normal','Body Text','First Paragraph']:
        if name in d.styles:
            st=d.styles[name];st.font.name='Cambria';st.font.size=Pt(11)
            st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=next(s for s in d.styles if s.name==name);st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(19 if name=='Title' else 12 if name=='Subtitle' else 13)
        st.paragraph_format.keep_with_next=True
    for p in d.paragraphs:
        p.paragraph_format.widow_control=True
        if p.text.startswith(('Table ','Figure ')): p.paragraph_format.keep_with_next=p.text.startswith('Table ')
    for tab in d.tables:
        tab.autofit=False
        for ri,row in enumerate(tab.rows):
            trpr=row._tr.get_or_add_trPr()
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for e in ['top','bottom','left','right']:
                    b=OxmlElement('w:'+e);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                pr.append(borders)
                if ri==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');pr.append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    for r in p.runs:r.font.size=Pt(10)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.core_properties.author='Wonsik Choi; Jeongin Choi'
    d.save(P/(stem+'.docx'))
if __name__=='__main__':build('manuscript.md','WRRA_Information_Weight_v0_3_EN')

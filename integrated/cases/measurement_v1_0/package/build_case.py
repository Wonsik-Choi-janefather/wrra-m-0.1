from pathlib import Path
import json, subprocess, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

R=Path(__file__).resolve().parent
OUT=R/'deliverables';OUT.mkdir(exist_ok=True)
res=json.loads((R/'expected/results.json').read_text())
PREFIX='WRRA_M_Integrated_Measurement_Case_v1_0'
DATE='2026_10_05'

EN=r'''# A Single Measurement in the WRRA M Integrated Model
## From address preparation to a finite record and homogeneous energy response

Wonsik Choi and Jeongin Choi

Corresponding author Wonsik Choi · Independent researcher Seoul Republic of Korea  
janefather@gmail.com · ORCID 0009-0001-4263-9772

Review case 1.0 · 5 October 2026  
Companion to the integrated manuscript DOI 10.5281/zenodo.23126800 [1]

# 1 Purpose and assessment

We give reviewers one executable case in the completed WRRA M integrated model. A finite address distribution prepares a proton internal state. An ideal write transfers its state and record costs to the same energy ledger used for homogeneous response.

**Verified inputs.** We retain the published address, internal Hamiltonian, constants and inherited energy calibration. The preparation strength, measurement basis, record gap, supply boundary and clock mapping are declared choices.

**WRRA transformation.** Address conditioning fixes an excited population. Its energy replaces the reference phenotype allocation. A supplier pays one write; the adopted volume law then gives pressure and deceleration from that energy.

**Outputs and falsifiers.** Baseline excitation is 34.4764357391 MeV. The mean supplier transfers 7.0378712840 × 10⁻¹² J and retains a positive balance. A failed state normalization, unmatched energy transfer, pressure derivative or proper-time transition rejects the corresponding connection.

Table 1 Fixed case inputs and declared roles

| Input | Value and role |
|---|---|
| Address support | $2\leq n\leq10^6$; finite source labels |
| Address parameters | $\alpha=1.8996876950554356$, $\beta=0.8654570124136961$ |
| Internal support | Two inherited modes within the symmetric 95D space |
| Preparation | $a_p=0.2$, selector prime 3, first excited mode |
| Energy fractions | $f_\phi=0.0493$, $f_D=0.265$, $f_R=0.6857$ |
| Reference density and volume | $u_0=7.6689477678\times10^{-10}\ \mathrm{J\,m^{-3}}$, $V_0=1\ \mathrm{m^3}$ |

We distinguish the individual write and controller from the expected-occupancy ledger. $D$ is the nonphenotypic resident sector. Prime-parts assembly rules are not used.

<<<PAGE>>>

# 2 Address conditioning and reference energy

Within the admitted phenotype branch, normalize the address weights as $w_n=|\phi_n|^2/\sum_m|\phi_m|^2$. Let $\nu_3(n)$ count the factors of 3 in address $n$, and let $\chi_n=1$ when that count is odd. The original address computation yields

$$t_3=\sum_n w_n\chi_n=0.3998832725128894,\qquad p_e=a_pt_3=0.07997665450257789.$$

The two selected modes $|g\rangle$ and $|e\rangle$ belong to the same inherited Hamiltonian $H$. We prepare the address-dephased internal mixture

$$\rho_0=(1-p_e)|g\rangle\langle g|+p_e|e\rangle\langle e|,\qquad \Delta E=431.0812443150\ \mathrm{MeV}.$$

Here $\Delta E$ is the first selected gap of the full 95D generator. Its two-mode subspace is invariant under that generator; it is not an unrelated replacement Hamiltonian. Thus $\operatorname{Tr}(\rho_0H)-E_g=p_e\Delta E=34.47643573912226\ \mathrm{MeV}$.

**Reference matching.** We adopt the inherited pure-proton illustration at $V_0$. The exact conversion is $1\ \mathrm{MeV}=1.602176634\times10^{-13}\ \mathrm{J}$. Set

$$\bar N=\frac{E_{\phi,\mathrm{ref}}}{E_{p,\mathrm{rest}}+p_e\Delta E}=0.2425893463937076,\qquad E_{\phi,\mathrm{ref}}=f_\phi u_0V_0.$$

Both denominator terms are expressed in J before evaluation. $\bar N$ is an externally matched expected occupancy. It is neither a derived cosmic proton abundance nor a fractional proton realization. The integer one-proton calculation on the next page has its own controller budget.

Table 2 Replacement of the inherited phenotype allocation

| Reference quantity | Energy in J |
|---|---:|
| Rest energy per proton | $1.5032776180\times10^{-10}$ |
| Expected rest allocation | $3.6467913480\times10^{-11}$ |
| Expected initial excitation | $1.3399990150\times10^{-12}$ |
| Rest plus excitation | $3.7807912495\times10^{-11}$ |
| Original phenotype allocation | $3.7807912495\times10^{-11}$ |

We replace the last row with the preceding split. Adding the split to the original total a second time would duplicate exactly $E_{\phi,\mathrm{ref}}$. The $D$ and $R$ allocations retain their inherited energies. Address branch probability 0.05 and energy fraction 0.0493 have different roles and are kept separate.

<<<PAGE>>>

# 3 One write and conditional time prediction

Choose $|\pm\rangle=(|g\rangle\pm|e\rangle)/\sqrt2$, with $P_\pm=|\pm\rangle\langle\pm|$. The blank record is $|0\rangle$ with gap $\varepsilon_r=3.5596112121\times10^{-15}\ \mathrm{J}$ from the electron-mode energy $\mu$ defined below. The original write coupling is

$$W=P_+\otimes I+P_-\otimes X,\qquad H_s=\operatorname{diag}(0,\Delta E),\quad H_r=\operatorname{diag}(0,\varepsilon_r).$$

$W$ is unitary, and the pointer maps the two outcomes to bits 0 and 1. Each conditional outcome has probability 1/2. Ignoring the outcome leaves system state $I/2$ and mean record energy $\varepsilon_r/2$. The relative system energy therefore changes from $p_e\Delta E$ to $\Delta E/2$.

For **one conditioned proton**, the mean write work is

$$w_{\mathrm{write}}=(\tfrac12-p_e)\Delta E+\tfrac12\varepsilon_r=2.9011460679\times10^{-11}\ \mathrm{J}.$$

All energies in this equation are evaluated in J. The individual original controller starts with $\Delta E+\varepsilon_r=6.9070389311\times10^{-11}\ \mathrm{J}$ and retains $4.0058928632\times10^{-11}\ \mathrm{J}$ after the mean write. It covers both outcome-dependent costs. This controller is an externally controlled budget; an autonomous energy-conserving apparatus Hamiltonian has not been constructed.

**Complete source routing.** Conditional measurement splits only the $\phi$ branch. The five probabilities are $0.025$, $0.025$, $0.268$, $0.007772944563$, and $0.674227055437$ for $\phi_+$, $\phi_-$, $D$, initial reflection and normal return. They sum to one; residents alone sum to 0.318.

**Proper-time prediction.** Adopt $\mu=510998.95069/23\ \mathrm{eV}$, $\omega=\mu/\hbar$ and $\delta\tau=0.1/\omega=2.9626039328317383\times10^{-21}\ \mathrm{s}$. We convert $\mu$ to J in the rate equation. After a selected first result,

$$U(\tau)=e^{-iH_s\tau/\hbar},\qquad P_{\mathrm{same}}(k)=\cos^2\!\left(\frac{\Delta E\,k\delta\tau}{2\hbar}\right).$$

| Separation in adopted frames | Same outcome probability |
|---:|---:|
| 1 | 0.675164206721 |
| 2 | 0.122729997265 |
| 4 | 0.569330619854 |

The clock mapping is an adopted construction, not an identified physical clock rate. These are conditional transition predictions. They do not include the energy of a repeated write/reset cycle or cumulative storage of all simulated records.

<<<PAGE>>>

# 4 The same energy in the homogeneous ledger

To connect the individual state calculation to the inherited reference volume, multiply its mean increments by $\bar N$. This defines a mean ensemble ledger for one ideal write per occupant. It does not make the small volume budget sufficient to perform one actual proton write; the individual controller on page 3 provides that separate illustration.

The mean internal increase is $7.0374395222\times10^{-12}\ \mathrm{J}$ and the mean record increase is $4.3176187869\times10^{-16}\ \mathrm{J}$. Their sum, $\delta E_{\mathrm{write}}=7.0378712840\times10^{-12}\ \mathrm{J}$, is supplied once. The finite mean supplier begins with $7.5615824991\times10^{-12}\ \mathrm{J}$.

Table 3 Mean energy accounts at the same reference volume

| Boundary quantity | Before | After |
|---|---:|---:|
| System including record in J | $7.6689477678\times10^{-10}$ | $7.7393264807\times10^{-10}$ |
| Supplier in J | $7.5615824991\times10^{-12}$ | $5.2371121503\times10^{-13}$ |
| System plus supplier in J | $7.7445635928\times10^{-10}$ | $7.7445635928\times10^{-10}$ |
| Pressure in Pa | $-5.2585974844\times10^{-10}$ | $-5.2585974844\times10^{-10}$ |
| $q$ with supplier outside | −0.528550000000 | −0.519196728075 |
| $q$ with pressureless supplier inside | −0.518507515893 | −0.518507515893 |

**Volume law and boundary.** At fixed comoving occupancy and fixed state, rest, excitation, $D$ and the modeled record are volume-independent energies. The $R$ sector is extensive in volume:

$$E_{\mathrm{out}}(V)=M+\delta E_{\mathrm{write}}+LV,\quad L=\frac{E_{R,0}}{V_0},\quad P=-\frac{\partial E}{\partial V}=-L.$$

$M$ contains the pre-write constant energies. Before the write, omit $\delta E_{\mathrm{write}}$. With $u=E/V$, the adopted homogeneous response is $q=\tfrac12(1+3P/u)$. The derivative and this response use the same energy expression.

If the supplier is outside, the system gains energy and its $q$ changes as shown. If a pressureless comoving supplier is included, its loss exactly offsets the system gain, so total energy and $q$ remain constant. The two initial $q$ values differ because the two boundaries contain different initial energies. The enlarged boundary includes the write supplier defined here, not an independently added copy of the individual controller budget.

These values are boundary diagnostics at $V_0$ under the stated volume law. They do not constitute a measured expansion change, a time-resolved source $Q(\tau)$, or a common SI covariant evolution of an apparatus and the cosmological background.

<<<PAGE>>>

# 5 Reproduction and review questions

We replayed unchanged bridge programs 0.2, 0.3, 0.5 and 0.6, forwarding the fresh state and basis. Their 95, 134, 134 and 87 checks pass, including controls beyond this case. The additional 26 contracts test this case's transfers and predictions. These are software consistency checks, not physical experiments.

Table 4 Controls that can reject or distinguish the construction

| Control | Calculated response or rejection |
|---|---|
| Change source phase by 0.7 rad | Address-dephased excitation unchanged within numerical tolerance |
| Double adopted frame duration | First-repeat probability changes from 0.675164206721 to 0.122729997265 |
| Reduce mean supplier to half the required work | Original exchange rule rejects the transfer |
| Add rest and excitation to the unreplaced phenotype energy | Duplicate $3.7807912495\times10^{-11}\ \mathrm{J}$ detected |
| Differentiate energy with respect to volume | Finite difference agrees with pressure $-L$ |

The package includes frozen source, inputs and hashes, executable code, expected results and CSV ledgers. `reference_boundary.json` retains the published comparison. Word equations are editable.

With Python and `requirements.txt`, run `python reproduce.py --out replay` in the unpacked package. It verifies the source hash and writes fresh results and logs. Success means 450 original checks, 26 case contracts and zero failures. Comparisons use declared numerical tolerances; displayed values are rounded.

**Questions for a reviewer**

1. Is the address selector a justified physical preparation of $\rho_0$, and do connected inputs affect the state as declared?
2. Are energy replacement and matched $\bar N$ legitimate mean calibration without implying a derived species abundance?
3. Do the instrument, supplier boundary and volume law support the calculated work and homogeneous response?
4. Which preparation, timing and record measurements distinguish these predictions from alternative rules at the same calibration?

**References**

[1] W. Choi and J. Choi. *WRRA M Integrated Upstream and Downstream Model*. Integrated 1.0 reviewed r1, 3 October 2026. https://doi.org/10.5281/zenodo.23126800

[2] W. Choi and J. Choi. *WRRA M Bridge 0.5–0.8 Reviewed r1*. Frozen executable source included unchanged in this package. Its source identity and the enclosing integrated package are recorded by SHA256 in `inputs.json`.
'''

KO=r'''# 상류 주소에서 측정 기록과 팽창 장부까지
## WRRA M 통합모형의 단일 측정 사례

최원식 Wonsik Choi · 최정인 Jeongin Choi

교신저자 최원식 · 독립연구자 서울 대한민국  
janefather@gmail.com · ORCID 0009-0001-4263-9772

검토용 사례 1.0 · 2026년 10월 5일  
통합논문 DOI 10.5281/zenodo.23126800의 동반 사례 [1]

# 1 목적과 판정 경로

완성된 WRRA M 상하류 통합모형에서 검토자가 따라갈 수 있는 사례 하나를 제시한다. 유한 주소분포가 양성자의 내부 상태를 준비하고, 이상적 기록 연산의 상태·기록 비용을 균질 응답에 사용하는 같은 에너지 장부로 넘긴다.

**검증 입력** 기존에 공개한 주소, 내부 해밀토니언, 상수와 상속 에너지 보정을 유지한다. 준비 강도, 측정 기저, 기록 에너지 간격, 공급 경계와 시계 사상은 명시한 구성 선택이다.

**WRRA 변환** 주소 조건화로 들뜬 상태의 점유율을 정하고 그 에너지로 기준 표현형 배정량을 대체한다. 공급원이 한 번의 기록 비용을 지불한다. 채택한 부피 법칙으로 같은 에너지의 압력과 감속계수를 계산한다.

**산출값과 반증조건** 기준 들뜸 에너지는 34.4764357391 MeV이다. 평균 공급원은 7.0378712840 × 10⁻¹² J를 전달하고 양의 잔액을 남긴다. 상태 정규화, 에너지 전달, 압력 미분 또는 고유시간 전이의 불일치는 해당 연결을 기각한다.

표 1 고정 입력과 구성 역할

| 입력 | 값과 역할 |
|---|---|
| 주소 지지집합 | $2\leq n\leq10^6$인 유한 원천 표지 |
| 주소 매개변수 | $\alpha=1.8996876950554356$, $\beta=0.8654570124136961$ |
| 내부 지지공간 | 대칭 95차원 공간에서 상속한 두 모드 |
| 상태 준비 | $a_p=0.2$, 선택 소수 3, 첫 들뜬 모드 |
| 에너지 분율 | $f_\phi=0.0493$, $f_D=0.265$, $f_R=0.6857$ |
| 기준 밀도와 부피 | $u_0=7.6689477678\times10^{-10}\ \mathrm{J\,m^{-3}}$, $V_0=1\ \mathrm{m^3}$ |

한 양성자의 기록·제어기 계산과 기대 점유수 장부를 구분한다. $D$는 비표현형 거주 부문이다. 부품 계열의 소수 조립 규칙은 사용하지 않는다.

<<<PAGE>>>

# 2 주소 조건화와 기준 에너지

통과한 표현형 가지 안에서 주소 가중치를 $w_n=|\phi_n|^2/\sum_m|\phi_m|^2$로 정규화한다. $\nu_3(n)$은 주소 $n$의 소인수 3의 개수이며, 그 개수가 홀수이면 $\chi_n=1$로 둔다. 원본 주소 계산은 다음 값을 준다.

$$t_3=\sum_n w_n\chi_n=0.3998832725128894,\qquad p_e=a_pt_3=0.07997665450257789.$$

선택한 두 모드 $|g\rangle$, $|e\rangle$는 같은 상속 해밀토니언 $H$의 상태이다. 주소를 탈위상한 내부 혼합 상태를 준비한다.

$$\rho_0=(1-p_e)|g\rangle\langle g|+p_e|e\rangle\langle e|,\qquad \Delta E=431.0812443150\ \mathrm{MeV}.$$

$\Delta E$는 전체 95차원 생성자에서 선택한 첫 에너지 간격이다. 두 모드의 지지공간은 그 생성자 아래에서 불변이다. 따라서 별개의 해밀토니언으로 교체하지 않고 $\operatorname{Tr}(\rho_0H)-E_g=p_e\Delta E=34.47643573912226\ \mathrm{MeV}$를 얻는다.

**기준량 맞춤** $V_0$에서 상속 보정의 순수 양성자 예시를 채택한다. 정확한 단위 변환은 $1\ \mathrm{MeV}=1.602176634\times10^{-13}\ \mathrm{J}$이다. 다음과 같이 정한다.

$$\bar N=\frac{E_{\phi,\mathrm{ref}}}{E_{p,\mathrm{rest}}+p_e\Delta E}=0.2425893463937076,\qquad E_{\phi,\mathrm{ref}}=f_\phi u_0V_0.$$

계산할 때 분모의 두 에너지를 모두 J로 변환한다. $\bar N$은 외부 기준에 맞춘 기대 점유수이다. 우주의 양성자 수를 유도하거나 양성자 일부를 실제로 측정한 값으로 해석하지 않는다. 다음 쪽의 양성자 한 개 계산에는 별도의 제어기 예산이 있다.

표 2 상속 표현형 배정량의 대체

| 기준 항목 | 에너지 J |
|---|---:|
| 양성자 한 개의 정지에너지 | $1.5032776180\times10^{-10}$ |
| 기대 정지에너지 배정량 | $3.6467913480\times10^{-11}$ |
| 기대 초기 들뜸 에너지 | $1.3399990150\times10^{-12}$ |
| 정지에너지와 들뜸의 합 | $3.7807912495\times10^{-11}$ |
| 기존 표현형 배정량 | $3.7807912495\times10^{-11}$ |

기존 배정량을 바로 위의 분할로 대체한다. 원래 총량에 분할 에너지를 다시 더하면 정확히 $E_{\phi,\mathrm{ref}}$가 중복된다. $D$와 $R$의 에너지는 상속 값을 유지한다. 주소 가지 확률 0.05와 에너지 분율 0.0493은 서로 다른 역할을 갖는다.

<<<PAGE>>>

# 3 한 번의 기록과 조건부 시간 예측

$|\pm\rangle=(|g\rangle\pm|e\rangle)/\sqrt2$와 $P_\pm=|\pm\rangle\langle\pm|$를 선택한다. 빈 기록은 $|0\rangle$이고 기록 간격 $\varepsilon_r=3.5596112121\times10^{-15}\ \mathrm{J}$는 아래에서 정의하는 전자 모드 에너지 $\mu$이다. 원본 기록 결합은 다음과 같다.

$$W=P_+\otimes I+P_-\otimes X,\qquad H_s=\operatorname{diag}(0,\Delta E),\quad H_r=\operatorname{diag}(0,\varepsilon_r).$$

$W$는 유니터리이며 포인터는 두 결과를 비트 0과 1에 대응시킨다. 각 조건부 결과의 확률은 1/2이다. 결과를 구분하지 않으면 내부 상태는 $I/2$, 평균 기록 에너지는 $\varepsilon_r/2$가 된다. 상대 내부 에너지는 $p_e\Delta E$에서 $\Delta E/2$로 바뀐다.

**조건부 양성자 한 개**에 대한 평균 기록 일은 다음과 같다.

$$w_{\mathrm{write}}=(\tfrac12-p_e)\Delta E+\tfrac12\varepsilon_r=2.9011460679\times10^{-11}\ \mathrm{J}.$$

이 식의 에너지는 모두 J로 계산한다. 원본의 개별 제어기는 $\Delta E+\varepsilon_r=6.9070389311\times10^{-11}\ \mathrm{J}$로 시작하며 평균 기록 후 $4.0058928632\times10^{-11}\ \mathrm{J}$를 남긴다. 두 결과 각각의 비용도 감당한다. 이는 외부 제어를 받는 예산이며, 자율적 에너지 보존 장치의 해밀토니언까지 구성한 것은 아니다.

**전체 원천 경로** 조건부 측정은 $\phi$ 가지만 나눈다. $\phi_+$, $\phi_-$, $D$, 초기 반사, 정상 반환의 확률은 각각 $0.025$, $0.025$, $0.268$, $0.007772944563$, $0.674227055437$이다. 합은 1이며 거주 부문의 합은 0.318이다.

**고유시간 예측** $\mu=510998.95069/23\ \mathrm{eV}$, $\omega=\mu/\hbar$, $\delta\tau=0.1/\omega=2.9626039328317383\times10^{-21}\ \mathrm{s}$를 채택한다. 각속도 식에서는 $\mu$를 J로 변환한다. 첫 결과를 선택한 뒤 다음을 계산한다.

$$U(\tau)=e^{-iH_s\tau/\hbar},\qquad P_{\mathrm{same}}(k)=\cos^2\!\left(\frac{\Delta E\,k\delta\tau}{2\hbar}\right).$$

| 채택한 프레임 간격 | 같은 결과가 나올 확률 |
|---:|---:|
| 1 | 0.675164206721 |
| 2 | 0.122729997265 |
| 4 | 0.569330619854 |

시계 사상은 채택한 구성이며 실제 시계 속도를 식별한 값은 아니다. 표는 조건부 전이 예측이다. 반복 기록·초기화 주기의 에너지와 모든 모의 기록을 누적 저장하는 에너지는 포함하지 않는다.

<<<PAGE>>>

# 4 같은 에너지의 균질 장부

개별 상태 계산을 상속 기준 부피에 연결하기 위해 평균 증가량에 $\bar N$을 곱한다. 이는 점유자마다 한 번의 이상적 기록을 적용하는 평균 앙상블 장부이다. 이 작은 부피의 예산으로 실제 양성자 한 개의 기록을 지불한다고 해석하지 않는다. 3쪽의 개별 제어기가 그 별도 계산을 제공한다.

평균 내부 증가량은 $7.0374395222\times10^{-12}\ \mathrm{J}$, 평균 기록 증가량은 $4.3176187869\times10^{-16}\ \mathrm{J}$이다. 합계 $\delta E_{\mathrm{write}}=7.0378712840\times10^{-12}\ \mathrm{J}$를 한 번만 공급한다. 평균 공급원의 초기 예산은 $7.5615824991\times10^{-12}\ \mathrm{J}$이다.

표 3 같은 기준 부피에서의 평균 에너지 장부

| 경계 항목 | 이전 | 이후 |
|---|---:|---:|
| 기록 포함 계의 에너지 J | $7.6689477678\times10^{-10}$ | $7.7393264807\times10^{-10}$ |
| 공급원 에너지 J | $7.5615824991\times10^{-12}$ | $5.2371121503\times10^{-13}$ |
| 계와 공급원의 합 J | $7.7445635928\times10^{-10}$ | $7.7445635928\times10^{-10}$ |
| 압력 Pa | $-5.2585974844\times10^{-10}$ | $-5.2585974844\times10^{-10}$ |
| 공급원을 밖에 둔 $q$ | −0.528550000000 | −0.519196728075 |
| 무압력 공급원을 포함한 $q$ | −0.518507515893 | −0.518507515893 |

**부피 법칙과 경계** 공이동 점유수와 상태를 고정하면 정지에너지, 들뜸, $D$, 모형 기록은 부피에 무관한 에너지이다. $R$ 부문의 에너지는 부피에 비례한다.

$$E_{\mathrm{out}}(V)=M+\delta E_{\mathrm{write}}+LV,\quad L=\frac{E_{R,0}}{V_0},\quad P=-\frac{\partial E}{\partial V}=-L.$$

$M$은 기록 이전의 상수 에너지들을 포함한다. 기록 이전에는 $\delta E_{\mathrm{write}}$를 뺀다. $u=E/V$로 두면 채택한 균질 응답은 $q=\tfrac12(1+3P/u)$이다. 미분과 응답이 같은 에너지 식을 사용한다.

공급원을 경계 밖에 두면 계의 에너지가 늘어나 표와 같이 $q$가 바뀐다. 무압력 공이동 공급원을 포함하면 그 감소량이 계의 증가량을 정확히 상쇄하여 총에너지와 $q$가 일정하다. 두 초기 $q$가 다른 이유는 두 경계의 초기 에너지가 다르기 때문이다. 확대 경계에는 여기서 정의한 기록 공급원만 포함하며 개별 제어기의 예산을 별도 사본으로 더하지 않는다.

이 값들은 명시한 부피 법칙 아래 $V_0$에서의 경계 진단이다. 실제 팽창 변화의 측정, 시간에 따른 공급항 $Q(\tau)$ 또는 장치와 우주 배경을 함께 진화시키는 SI 공변 계산까지 수행한 결과는 아니다.

<<<PAGE>>>

# 5 재현과 검토 질문

원본 브리지 0.2, 0.3, 0.5, 0.6을 그대로 재실행하고 새 상태와 기저를 전달했다. 이 사례 밖의 대조를 포함한 점검 95, 134, 134, 87개가 통과한다. 추가 점검 26개는 사례의 전달과 예측을 검사한다. 모두 소프트웨어 일관성 점검이며 물리 실험 횟수는 아니다.

표 4 구성을 구별하거나 기각하는 대조

| 대조 | 계산 응답 또는 거부 |
|---|---|
| 원천 위상을 0.7 rad 바꿈 | 주소 탈위상 들뜸 에너지는 수치 허용오차 안에서 동일 |
| 채택한 프레임 시간을 두 배로 함 | 첫 반복 확률이 0.675164206721에서 0.122729997265로 변함 |
| 평균 공급을 필요 일의 절반으로 축소 | 원본 교환 규칙이 전달을 거부 |
| 대체하지 않은 표현형 에너지에 정지·들뜸을 추가 | $3.7807912495\times10^{-11}\ \mathrm{J}$ 중복 검출 |
| 에너지를 부피로 미분 | 유한 차분이 압력 $-L$과 일치 |

패키지에는 고정 소스, 입력과 해시, 실행 코드, 기대 결과와 CSV 장부가 있다. `reference_boundary.json`은 기존 공개 비교값이다. Word 수식은 편집할 수 있다.

Python과 `requirements.txt`를 준비하고 압축을 푼 폴더에서 `python reproduce.py --out replay`를 실행한다. 소스 해시를 검사하고 새 결과와 로그를 만든다. 정상 결과는 원본 점검 450개, 사례 점검 26개, 실패 0개이다. 명시한 수치 허용오차로 비교하며 본문 값은 반올림했다.

**검토자에게 요청하는 질문**

1. 주소 선택자가 $\rho_0$의 물리적 준비를 충분히 설명하는가? 연결된 입력이 선언대로 상태에 영향을 주는가?
2. 에너지 대체와 맞춘 $\bar N$이 종별 존재량 유도를 주장하지 않는 평균 보정으로 적절한가?
3. 측정, 공급 경계와 부피 법칙이 계산한 일과 균질 응답을 지지하는가?
4. 같은 보정의 다른 규칙과 예측을 구별할 상태 준비, 시간 설정과 기록 측정은 무엇인가?

**참고문헌**

[1] W. Choi and J. Choi. *WRRA M Integrated Upstream and Downstream Model*. Integrated 1.0 reviewed r1, 2026년 10월 3일. https://doi.org/10.5281/zenodo.23126800

[2] W. Choi and J. Choi. *WRRA M Bridge 0.5–0.8 Reviewed r1*. 이 패키지에 원본 실행 소스를 변경 없이 포함했다. 소스와 이를 담은 통합 패키지의 SHA256은 `inputs.json`에 기록했다.
'''

def style_doc(path,lang):
    d=Document(path);sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
    sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
    font='NanumGothic' if lang=='KO' else 'Calibri'
    for style in d.styles:
        if style.type==1:
            style.font.name=font;style.font.color.rgb=RGBColor(0,0,0)
            if style._element.rPr is not None:
                for b in style._element.rPr.findall(qn('w:bdr')):style._element.rPr.remove(b)
            if style._element.pPr is not None:
                for b in style._element.pPr.findall(qn('w:pBdr')):style._element.pPr.remove(b)
            rpr=style.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
            if rf is None:rf=OxmlElement('w:rFonts');rpr.append(rf)
            for key in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+key),font)
    normal=d.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.14
    for name in ('Body Text','First Paragraph','Compact'):
        if name in d.styles:
            d.styles[name].base_style=normal;d.styles[name].font.size=Pt(11)
            d.styles[name].paragraph_format.space_after=Pt(6)
            d.styles[name].paragraph_format.line_spacing=1.14
    d.styles['Title'].font.size=Pt(19);d.styles['Title'].font.bold=False
    d.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Title'].paragraph_format.space_before=Pt(0)
    d.styles['Title'].paragraph_format.space_after=Pt(7)
    d.styles['Subtitle'].font.size=Pt(12);d.styles['Subtitle'].font.italic=False
    d.styles['Subtitle'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    d.styles['Subtitle'].paragraph_format.space_after=Pt(9)
    for st in d.styles:
        if st.style_id not in ('Heading1','Heading2','Heading3'):continue
        st.font.size=Pt(13 if st.style_id=='Heading1' else 12)
        st.font.bold=True
        st.paragraph_format.space_before=Pt(9)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.keep_with_next=True
    d.paragraphs[0].style='Title';d.paragraphs[1].style='Subtitle'
    for idx,p in enumerate(d.paragraphs):
        if p.text=='CASEPAGEBREAK':
            p.clear();p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
            p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1
        elif idx in (2,3,4):
            for r in p.runs:r.font.size=Pt(9.5)
            p.paragraph_format.space_after=Pt(4)
        if p._p.findall('.//'+qn('m:oMathPara')):
            p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.line_spacing=1.05
    for ti,t in enumerate(d.tables):
        t.autofit=False
        widths=[2.05,5.05] if len(t.columns)==2 and ti in (0,4) else ([4.0,3.1] if len(t.columns)==2 else [2.65,2.225,2.225])
        if ti==2:widths=[3.6,3.5]
        if ti==4:widths=[2.45,4.65]
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        pr=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        pr.append(borders)
        for ri,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();n=OxmlElement('w:cantSplit');trpr.append(n)
            if ri==0:
                hdr=OxmlElement('w:tblHeader');trpr.append(hdr)
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);tcp=cell._tc.get_or_add_tcPr()
                mar=OxmlElement('w:tcMar')
                for edge in ('top','bottom','left','right'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65' if edge in ('top','bottom') else '90');e.set(qn('w:type'),'dxa');mar.append(e)
                tcp.append(mar)
                if ri==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');tcp.append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.06
                    for r in p.runs:r.font.size=Pt(10.5);r.bold=True if ri==0 else r.bold
    for p in d.paragraphs:
        if p.text.startswith(('Table ','표 ')):
            p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(10.5);r.italic=False
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run('WRRA M Integrated Case 1.0 | ');run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.core_properties.title='상류 주소에서 측정 기록과 팽창 장부까지' if lang=='KO' else 'A Single Measurement in the WRRA M Integrated Model'
    d.core_properties.author='Wonsik Choi; Jeongin Choi'
    d.core_properties.subject='Conditional measurement and energy boundary case for review'
    # Use an installed OpenType math font and explicit normal text for upright units.
    mathfont=d.settings.element.find('.//'+qn('m:mathFont'))
    if mathfont is not None:mathfont.set(qn('m:val'),'Latin Modern Math')
    for rpr in d.element.findall('.//'+qn('m:rPr')):
        sty=rpr.find(qn('m:sty'))
        if sty is not None and sty.get(qn('m:val'))=='p':
            rpr.remove(sty)
            nor=OxmlElement('m:nor');nor.set(qn('m:val'),'1');rpr.append(nor)
    # Keep literal bars in bra/ket and modulus expressions: LibreOffice otherwise
    # substitutes some closing delimiter glyphs during OMML-to-PDF conversion.
    def literal(token):
        rr=OxmlElement('m:r');rp=OxmlElement('m:rPr');no=OxmlElement('m:nor');no.set(qn('m:val'),'1');rp.append(no);rr.append(rp)
        tt=OxmlElement('m:t');tt.text=token;rr.append(tt);return rr
    for de in list(d.element.findall('.//'+qn('m:d'))):
        pr=de.find(qn('m:dPr'))
        if pr is None:continue
        beg=pr.find(qn('m:begChr'));end=pr.find(qn('m:endChr'))
        bv=beg.get(qn('m:val')) if beg is not None else '('
        ev=end.get(qn('m:val')) if end is not None else ')'
        if bv!='|' and ev!='|':continue
        parent=de.getparent();index=parent.index(de);items=[literal(bv)]
        for e in de.findall(qn('m:e')):items.extend(list(e))
        items.append(literal(ev));parent.remove(de)
        for offset,item in enumerate(items):parent.insert(index+offset,item)
    for parent in d.element.iter():
        previous=None
        for child in list(parent):
            if child.tag!=qn('m:r'):previous=None;continue
            text=child.find(qn('m:t'));rp=child.find(qn('m:rPr'))
            nor=rp.find(qn('m:nor')) if rp is not None else None
            if text is None or not (text.text or '').isalpha() or nor is None:previous=None;continue
            if previous is not None:
                previous.find(qn('m:t')).text+=text.text;parent.remove(child)
            else:previous=child
    d.save(path)

for lang,text in [('KO',KO),('EN',EN)]:
    base=OUT/f'{PREFIX}_{lang}_{DATE}'
    base.with_suffix('.md').write_text(text.replace('<<<PAGE>>>','CASEPAGEBREAK').replace(r'\bar N',r'\overline{N}'),encoding='utf-8')
    subprocess.run(['pandoc',str(base.with_suffix('.md')),'--from','markdown+tex_math_dollars','--to','docx','--output',str(base.with_suffix('.docx'))],check=True)
    style_doc(base.with_suffix('.docx'),lang)
    print(base.with_suffix('.docx'))

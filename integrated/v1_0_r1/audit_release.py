"""Portable audit of actual bilingual manuscript values and frozen evidence.

Run from any directory: python /path/to/extracted/package/audit_release.py
Requires python-docx, lxml and pypdf. This is a computational/document audit.
"""
from pathlib import Path
import hashlib, json, math, re
from docx import Document
from lxml import etree
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
SUPER = str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻', '0123456789-')

def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p): return json.loads(p.read_text())
def numbers(text):
    text = text.replace('−', '-').replace('\\times', '×')
    text = re.sub(r'([⁰¹²³⁴⁵⁶⁷⁸⁹⁻]+)', lambda m: '^' + m[0].translate(SUPER), text)
    text = re.sub(r'10\^\{(-?\d+)\}', r'10^\1', text)
    pattern = r'-?\d+(?:\.\d+)?(?:\s*×\s*10\^?\s*-?\d+|[eE][+-]?\d+)?'
    values=[]
    for match in re.finditer(pattern, text):
        value=match[0]
        if '×' in value:
            a,b=re.split(r'\s*×\s*10\^?\s*',value)
            values.append(float(a)*10**int(b))
        else: values.append(float(value))
    return values

def audit():
    docs={l:Document(ROOT/f'manuscripts/WRRA_M_Integrated_1_0_{l}_2026_10_03.docx') for l in ('KO','EN')}
    equations={l:d._element.xpath('.//m:oMathPara/m:oMath') for l,d in docs.items()}
    assert all(len(x)==25 for x in equations.values())
    assert [etree.tostring(x,method='c14n') for x in equations['KO']]==[etree.tostring(x,method='c14n') for x in equations['EN']]
    assert all([len(t.rows) for t in d.tables]==[8,16,3,8,7,5,11,8] for d in docs.values())
    registered=read(ROOT/'numerical_claims.json')['claims']
    table_checks=prose_checks=0
    for c in registered:
        value=read(ROOT/c['source_file'])
        for key in c['json_path']: value=value[key]
        for lang,location in c['document_locations'].items():
            if location['type']=='table':
                text=docs[lang].tables[location['table_index']].cell(location['row_index'],location['column_index']).text
                actual=numbers(text)[location['number_index']]
                table_checks+=1
            else:
                lines=(ROOT/f'manuscripts/WRRA_M_Integrated_1_0_{lang}_2026_10_03.md').read_text().splitlines()
                text=lines[location['line_index']]
                actual=numbers(text)[location['number_index']]
                prose_checks+=1
            assert math.isclose(actual,float(value),rel_tol=c['relative_tolerance'],abs_tol=0), (lang,c['name'],actual,value,text)
    bridge=read(ROOT/'evidence/bridge/results.json')
    assert bridge['stage_case_checks']==1116 and bridge['review_passed']==188 and bridge['review_failed']==0
    assert all(row['failed']==0 for row in bridge['stages'])
    counts={}
    for n in (10,11,12):
        v=read(ROOT/f'evidence/downstream/calculations/wrra_m_0_{n}/results/verification.json')
        assert v['passed'] and all(c['passed'] for c in v['checks'])
        counts[str(n)]=v['check_count']
    assert counts=={'10':32,'11':28,'12':36}
    expected=read(ROOT/'evidence/downstream/replay_expected.json')
    for path,digest in expected.items(): assert sha(ROOT/'evidence/downstream'/path)==digest,path
    pdfs={}
    for lang in ('KO','EN'):
        pages=PdfReader(ROOT/f'manuscripts/WRRA_M_Integrated_1_0_{lang}_2026_10_03.pdf').pages
        assert all(len(p.extract_text().strip())>100 and '\ufffd' not in p.extract_text() for p in pages)
        pdfs[lang]=len(pages)
    manifest_count=0
    if (ROOT/'SHA256SUMS').exists():
        for line in (ROOT/'SHA256SUMS').read_text().splitlines():
            digest,name=line.split(maxsplit=1)
            assert sha(ROOT/name.lstrip('*'))==digest,name
            manifest_count+=1
    return dict(passed=True,edition='1.0-r1',numerical_claims=len(registered),actual_table_value_checks=table_checks,actual_prose_value_checks=prose_checks,bilingual_native_equations_identical=True,native_equations_per_language=25,tables_per_language=8,bridge_stage_checks=1116,bridge_review_checks=188,bridge_failed=0,downstream_distinct_checks=counts,downstream_distinct_total=96,downstream_identical_files=len(expected),pdf_pages=pdfs,manifest_entries=manifest_count,scope='Computational reproduction and document validation, not observational experiments or independent physical proofs.')

if __name__=='__main__': print(json.dumps(audit(),ensure_ascii=False,indent=2))

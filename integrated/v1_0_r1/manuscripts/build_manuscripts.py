from pathlib import Path
import re, subprocess, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parent

def setfont(style, latin, east, size):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for key, value in [('ascii',latin),('hAnsi',latin),('eastAsia',east),('cs',latin)]:
        rf.set(qn('w:'+key),value)

def equation_layout(text):
    def fix(m):
        s=m.group(1).strip()
        s=re.sub(r'\\rm\s+([A-Za-z]+)',r'\\mathrm{\1}',s)
        if '\\qquad\n' in s:
            s='\\begin{gathered}\n'+s.replace('\\qquad\n','\\\\\n')+'\n\\end{gathered}'
        return '$$'+s+'$$'
    return re.sub(r'\$\$([\s\S]*?)\$\$',fix,text)

for lang in ('KO','EN'):
    stem=f'WRRA_M_Integrated_1_0_{lang}_2026_10_03'
    src=ROOT/(stem+'.md')
    text=src.read_text()
    if lang=='KO':
        text=text.replace('표 7 apparent 충돌의 판별과 해결','표 7 겉보기 충돌의 판별과 해결').replace('durable 기록 매체','영속 기록 매체')
        if 'θj는' not in text:
            text=text.replace('위상만 읽어 가지를 버리면 원래 에너지를 복원할 수 없다.','위상만 읽어 가지를 버리면 원래 에너지를 복원할 수 없다. 식 20의 θj는 U† 고유값의 주값 위상으로 정의하여 양의 생성자 위상과 부호를 맞춘다.')
    else:
        if 'θj is the principal phase' not in text:
            text=text.replace('Reading phase alone and discarding branches loses the original energy.','Reading phase alone and discarding branches loses the original energy. In Eq. 20, θj is the principal phase of the eigenvalue of U†, matching the positive generator-phase convention.')
    src.write_text(text)
    temp=ROOT/(stem+'.typeset.md'); temp.write_text(equation_layout(text))
    target=ROOT/(stem+'.docx')
    subprocess.run(['pandoc',str(temp),'-o',str(target),'--from','markdown+tex_math_dollars','--standalone'],check=True)
    d=Document(target)
    latin='DejaVu Serif'; east='NanumMyeongjo'
    setfont(d.styles['Normal'],latin,east,10.5 if lang=='KO' else 10.4)
    pf=d.styles['Normal'].paragraph_format
    pf.line_spacing=1.16;pf.space_after=Pt(5.5);pf.widow_control=True
    for sn,size in [('Title',22),('Subtitle',12),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        if sn not in d.styles:
            d.styles.add_style(sn,WD_STYLE_TYPE.PARAGRAPH)
        setfont(d.styles[sn],latin,'NanumGothic',size)
        d.styles[sn].font.bold=sn!='Subtitle'
        d.styles[sn].paragraph_format.keep_with_next=True
        d.styles[sn].paragraph_format.space_before=Pt(11 if sn.startswith('Heading') else 0)
        d.styles[sn].paragraph_format.space_after=Pt(6)
    for sn in ('Source Code','Verbatim Char'):
        if sn in d.styles:
            setfont(d.styles[sn],'DejaVu Sans Mono','NanumGothic',8.5)
    d.paragraphs[0].style='Title';d.paragraphs[1].style='Subtitle'
    for idx,p in enumerate(d.paragraphs):
        if p.style.name in ('Heading1','Heading2','Heading3'):
            p.style='Heading '+p.style.name[-1]
        if idx<7:
            p.paragraph_format.keep_with_next=True
        if p._p.findall('.//'+qn('m:oMathPara')):
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(7)
            p.paragraph_format.keep_together=True
        if re.match(r'^(표 [1-7]|Table [1-7]) ',p.text):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(7)
            for run in p.runs:run.bold=True;run.font.size=Pt(9.4)
        if p.text.startswith('['):
            p.paragraph_format.space_after=Pt(5)
            for run in p.runs:run.font.size=Pt(9)
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
        sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
        sec.footer_distance=Cm(.85)
        p=sec.footer.paragraphs[0];p.alignment=1
        run=p.add_run();run.font.size=Pt(8)
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    for table in d.tables:
        table.autofit=False
        cols=len(table.columns)
        widths={2:[5.5,11.3],3:[3.8,6.5,6.5],4:[3.5,4.1,5.3,3.9]}[cols]
        for rowidx,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            ns=OxmlElement('w:cantSplit');trpr.append(ns)
            if rowidx==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for i,c in enumerate(row.cells):
                c.width=Cm(widths[i])
                c.vertical_alignment=1
                tcpr=c._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for key,val in [('top',75),('bottom',75),('left',100),('right',100)]:
                    el=OxmlElement('w:'+key);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
                tcpr.append(margins)
                if rowidx==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');tcpr.append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
                    for run in p.runs:
                        run.font.size=Pt(8.5 if lang=='EN' else 9)
                        run.bold=rowidx==0
    # No automatic horizontal rules on title or any heading.
    for p in d.paragraphs:
        pr=p._p.find(qn('w:pPr'))
        if pr is not None:
            for el in list(pr):
                if el.tag==qn('w:pBdr'):pr.remove(el)
    # Keep the short authors/availability section on one page.
    author_start=next(i for i,p in enumerate(d.paragraphs) if p.text in ('Authors and data availability','저자와 자료 공개'))
    for p in d.paragraphs[author_start:]:
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=p is not d.paragraphs[-1]
    d.paragraphs[-1].paragraph_format.keep_with_next=False
    d.core_properties.title=d.paragraphs[0].text
    d.core_properties.author='Wonsik Choi; Jeongin Choi'
    d.core_properties.subject='Finite upstream downstream bridge and common physical ledger'
    d.core_properties.keywords='WRRA M, minimum computation, energy, currents, measurement'
    d.save(target)
    temp.unlink()
    print(target.name)

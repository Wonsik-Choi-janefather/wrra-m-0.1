"""Audit published calculation evidence and the bilingual integrated manuscripts."""
from pathlib import Path
import hashlib, json, math, re, zipfile
from lxml import etree
from docx import Document
from pypdf import PdfReader
import argparse

HERE = Path(__file__).resolve().parent
WORK = HERE.parent
BASE = WORK / 'baseline_series'
REPLAY = WORK / 'replay_series'
ARCHIVE = Path('/workspace/scratch/WRRA_M_0_12_r1_Series_Release.zip')

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def load_result(n):
    return json.loads((BASE / f'calculations/wrra_m_0_{n}/results/results.json').read_text())

def cell_text(cell):
    return ''.join(cell._tc.xpath('.//w:t/text() | .//m:t/text()'))

def numeric(text):
    # Superscript exponents are presentation only; restore them for numerical comparison.
    supers = '⁰¹²³⁴⁵⁶⁷⁸⁹⁻'
    text = re.sub('([' + supers + ']+)',
                  lambda m: '^' + m.group().translate(str.maketrans(supers, '0123456789-')), text)
    text=text.replace('−','-')
    m = re.search(r'(-?\d+(?:\.\d+)?)\s*×\s*10\^?(-?\d+)', text)
    if m:
        return float(m[1]) * 10 ** int(m[2])
    return float(re.search(r'-?\d+(?:\.\d+)?', text)[0])

def near(actual, expected, abs_tol=0.0):
    assert math.isclose(actual, expected, rel_tol=5e-10, abs_tol=abs_tol), (actual, expected)

def audit():
    manifest_entries = 0
    for line in (BASE / 'SHA256SUMS').read_text().splitlines():
        expected, name = line.split(maxsplit=1)
        assert sha(BASE / name.lstrip('*')) == expected, name
        manifest_entries += 1
    replay_expected = json.loads((WORK / 'replay_expected.json').read_text())
    for name, expected in replay_expected.items():
        assert sha(REPLAY / name) == expected, name
        assert sha(BASE / name) == expected, name

    checks = {}
    for n in (10, 11, 12):
        d = json.loads((REPLAY / f'calculations/wrra_m_0_{n}/results/verification.json').read_text())
        assert d['passed'] and all(c['passed'] for c in d['checks'])
        checks[str(n)] = d['check_count']
    assert checks == {'10': 32, '11': 28, '12': 36}
    review = json.loads((REPLAY / 'review/0_10_to_0_12_r1/review_checks.json').read_text())
    assert review['passed'] and review['review_check_count'] == 9

    docs = {l: Document(HERE / f'WRRA_M_1_0_Downstream_Synthesis_{l}.docx') for l in ('EN','KO')}
    display = {l: d._element.xpath('.//m:oMathPara/m:oMath') for l,d in docs.items()}
    sig = lambda xs: [etree.tostring(x, method='c14n') for x in xs]
    assert len(display['EN']) == len(display['KO']) == 22
    assert sig(display['EN']) == sig(display['KO'])
    assert all(len(d.tables) == 6 for d in docs.values())
    assert all([len(t.rows) for t in d.tables] == [9,13,7,7,7,12] for d in docs.values())

    # The table schema and adopted/calculated cells must agree across translations.
    compared_cells = 0
    for ti, columns in [(0,[1]),(1,[0]),(2,[1,2]),(3,[1]),(4,[1,2,3])]:
        for ri in range(1,len(docs['EN'].tables[ti].rows)):
            for ci in columns:
                en = cell_text(docs['EN'].tables[ti].cell(ri,ci))
                ko = cell_text(docs['KO'].tables[ti].cell(ri,ci))
                if ti==4 and ri<=2 and ci==3:
                    assert (en,ko) == ('Not used','사용 안 함')
                else:
                    assert en == ko, (ti,ri,ci,en,ko)
                compared_cells += 1

    d10,d11,d12 = (load_result(n) for n in (10,11,12))
    for lang,doc in docs.items():
        assert cell_text(doc.tables[0].cell(7,1)).endswith('J')
        for idx in (5,10):
            eq=display[lang][idx]
            text=''.join(eq.xpath('.//m:t/text()',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}))
            assert '(' in text and ')' in text
        clock=''.join(display[lang][20].xpath('.//m:t/text()',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}))
        assert 'τ' in clock and 'ℏ' in clock and clock.count('exp')==2
        t3 = doc.tables[2]
        for ci,ref in [(1,d11['reference']),(2,d10['alternate_reference'])]:
            fractions=[float(x) for x in cell_text(t3.cell(1,ci)).split('/')]
            expected=[ref['sector_energy_fractions'][s] for s in ('phenotype','resident_nonphenotype','return')]
            for a,b in zip(fractions,expected):near(a,b)
            values=[ref['total_density_J_m3'],ref['total_pressure_Pa'],ref['deceleration_q'],
                    ref['local_readout']['v_total_km_s'],
                    ref['local_readout']['finite_patch_lensing']['alpha_patch_arcsec']]
            for ri,value in enumerate(values,2):near(numeric(cell_text(t3.cell(ri,ci))),value,abs_tol=5e-10 if ri==6 else 0)
        t4=doc.tables[3]
        units=d12['units']
        values=[units['mu_E_eV'],units['t_mu_s'],0.05,units['delta_tau_s'],128,d12['cavity_modes'][0]['length_m']]
        for ri,value in enumerate(values,1):near(numeric(cell_text(t4.cell(ri,1))),value)
        t5=doc.tables[4]
        expected=[]
        for eta in (0,0.5):
            row=next(r for r in d12['fold_modes'] if r['boundary_phase']==eta and r['mode']==23)
            expected.append((23,row['energy_eV'],None))
        for boundary,mode in [('dirichlet',1),('dirichlet',2),('periodic',0),('periodic',1)]:
            row=next(r for r in d12['cavity_modes'] if r['boundary']==boundary and r['mode']==mode)
            expected.append((mode,row['energy_eV'],row['kinetic_energy_eV']))
        for ri,values in enumerate(expected,1):
            for ci,value in enumerate(values,1):
                if value is not None:near(numeric(cell_text(t5.cell(ri,ci))),value,abs_tol=5e-10)
        source=(HERE / f'WRRA_M_1_0_Downstream_Synthesis_{lang}.md').read_text()
        assert f"{d12['legacy_clock_comparison']['old_frequency_over_SI_frequency']:.15e}".split('e')[0] in source
        assert d11['input_hash_sha256'] in source and d12['input_hash_sha256'] in source
        assert '0.13' in source and '0.17' in source and '0.12' in source
        assert 'Wonsik Choi' in source and 'Jeongin Choi' in source
        assert ('Unimplemented' if lang=='EN' else '미구현') in source

    pdf_pages={l:len(PdfReader(HERE / f'WRRA_M_1_0_Downstream_Synthesis_{l}.pdf').pages) for l in ('EN','KO')}
    return {
        'edition':'WRRA M 1.0 downstream consolidation of 0.1 through 0.12',
        'audit_date_utc':'2026-10-03', 'passed':True,
        'baseline_record':'https://zenodo.org/records/23113101',
        'baseline_archive_sha256':sha(ARCHIVE),
        'archive_current_manifest_entries_verified':manifest_entries,
        'clean_replay_files_identical':len(replay_expected),
        'distinct_component_check_groups':checks,
        'distinct_component_total':sum(checks.values()),
        'cross_version_review_checks':review['review_check_count'],
        'bilingual_identical_native_display_equations_per_edition':22,
        'tables_per_edition':6,'table_rows_with_headers':[9,13,7,7,7,12],
        'bilingual_table_cells_compared':compared_cells,
        'physical_calibration_electron_units_and_spectral_table_values_checked_against_raw_results':True,
        'pdf_pages':pdf_pages,
        'frozen_input_hashes':{str(n):load_result(n)['input_hash_sha256'] for n in (10,11,12)},
        'endpoint_commit':'b4aa53a481608add2f9b41e432ea516f6bc08fcb',
        'notes':[
            'The root SHA256SUMS is the current whole-series manifest. Version-specific manifests remain historical; SHA256SUMS_0_12 has two historical root metadata hashes that differ from the reviewed series metadata.',
            'Old document_checks.json files were retained as evidence for already published documents, not treated as new checks of this synthesis.',
            'The new synthesis was visually inspected separately in both languages after final rendering.',
            'This content audit precedes publication of the final 1.0 edition; publication is verified separately.',
        ],
    }

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--evidence-dir',type=Path,default=WORK)
    parser.add_argument('--archive',type=Path,default=WORK / 'WRRA_M_0_12_r1_Series_Release.zip')
    args=parser.parse_args()
    WORK=args.evidence_dir.resolve();BASE=WORK / 'baseline_series';REPLAY=WORK / 'replay_series';ARCHIVE=args.archive.resolve()
    result=audit()
    (HERE / 'WRRA_M_1_0_Validation_Report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))

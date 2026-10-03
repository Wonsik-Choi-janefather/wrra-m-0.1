from pathlib import Path
import re, subprocess
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
names={l:ROOT/f'WRRA_M_1_0_Downstream_Synthesis_{l}.md' for l in ('EN','KO')}
en=names['EN'].read_text();ko=names['KO'].read_text()
def polish_math_text(s):
 # Keep display mathematics editable; replace prose approximations with native inline math.
 for old,new in {
  'AD = $K_c$/2':r'$A_D=K_c/2$',
  'AR = $K_b$/2':r'$A_R=K_b/2$',
  'E/(u*V₀) = 1':r'$E/(u_*V_0)=1$',
  'P/u* = -0.682':r'$P/u_*=-0.682$',
  'E*/ℏ':r'$E_*/\hbar$',
 }.items():s=s.replace(old,new)
 expressions={
  'g = gM/[1 - exp(-sqrt(gM/aT))]':r'g=\frac{g_M}{1-\exp[-\sqrt{g_M/a_T}]}',
  'Qh = ζT κh²':r'Q_h=\zeta_T\kappa_h^2',
  '1/sqrt(r)':r'1/\sqrt r',
  'Φ(r) = -∫r^R g(s) ds':r'\Phi(r)=-\int_r^R g(s)\,ds',
  'E(p) = sqrt(Ee² + c²p²)':r'E(p)=\sqrt{E_e^2+c^2p^2}',
  'A = 1 + 2Φ/c²':r'A=1+2\Phi/c^2',
  'B = 1 - 2Φ/c²':r'B=1-2\Phi/c^2',
  'ωinfo = 0.7H0':r'\omega_{\mathrm{info}}=0.7H_0',
  'E* = ucrit V0':r'E_*=u_{\mathrm{crit}}V_0',
  'ΔL = ΔR = 0.04669193039':r'\Delta_L=\Delta_R=0.04669193039',
  'Ω(n) - 1':r'\Omega(n)-1',
  'K_c':r'K_c','K_b':r'K_b',
  'νφ = νD = 0':r'\nu_\varphi=\nu_D=0','νR = 3':r'\nu_R=3',
  'Aφ = I':r'A_\varphi=I','AD = Kc/2':r'A_D=K_c/2','AR = Kb/2':r'A_R=K_b/2',
  'λφ':r'\lambda_\varphi','λD':r'\lambda_D','λR':r'\lambda_R',
  'μE':r'\mu_E','tμ':r't_\mu','ετ':r'\epsilon_\tau','Δτ*':r'\Delta\tau_*','ℓμ':r'\ell_\mu',
  'ηB':r'\eta_B','ηs':r'\eta_s',
 }
 supers=str.maketrans('0123456789-','⁰¹²³⁴⁵⁶⁷⁸⁹⁻')
 parts=re.split(r'(\$\$.*?\$\$|\$[^$]*\$)',s,flags=re.S)
 for i,part in enumerate(parts):
  if part.startswith('$'):continue
  for old,new in expressions.items():part=part.replace(old,'$'+new+'$')
  subparts=re.split(r'(\$[^$]*\$)',part)
  for j,q in enumerate(subparts):
   if q.startswith('$'):continue
   q=re.sub(r'\^(-?\d+)',lambda m:m.group(1).translate(supers),q)
   for old,new in [('H0','H₀'),('q0','q₀'),('Ee','Eₑ'),('V0','V₀'),('Kc','K_c'),('Kb','K_b')]:q=q.replace(old,new)
   subparts[j]=q
  parts[i]=''.join(subparts)
 return ''.join(parts)
en=polish_math_text(en);ko=polish_math_text(ko)
en=en.replace('Correspondence Wonsik Choi Independent Researcher Seoul Republic of Korea','Correspondence: Wonsik Choi, Independent Researcher, Seoul, Republic of Korea')
en=en.replace(r'\downarrow SU(3)',r'\downarrow\mathrm{SU}(3)')
ko=ko.replace(r'\downarrow SU(3)',r'\downarrow\mathrm{SU}(3)')
eqs=re.findall(r'\$\$(.*?)\$\$',en,re.S)
replacements={
9:r'\begin{aligned}\sigma&=\sum_nw_n|n\rangle\langle n|\otimes\rho,\\g_s(n)&=1+\lambda_s\frac{\log n}{\log N},\qquad \mu_s=\sum_nw_ne_s(n)g_s(n).\end{aligned}',
11:r'\begin{aligned}\widehat E(a)&=V_0\sum_s\eta_s\mu_sa^{\nu_s}A_s,\\E_s&=V_0\eta_s\mu_sa^{\nu_s}\operatorname{Tr}(\rho A_s),\qquad E=\sum_sE_s.\end{aligned}',
13:r'\begin{aligned}H^2&=\frac{8\pi G}{3c^2}u,\qquad q=\frac{u+3P}{2u},\\a_T&=cH_0\sqrt{\frac{u_D}{8u_{\mathrm{crit}}}},\qquad \zeta_T\kappa_h^2=\frac{16\pi G}{c^4}(u_D+u_R).\end{aligned}',
14:r'\begin{aligned}U(\tau)&=e^{-i\widehat E\tau/\hbar},\qquad \rho(\tau)=U\rho_0U^\dagger,\\\operatorname{Tr}(\widehat E\dot\rho)&=0,\qquad \dot u+3H(u+P)=0.\end{aligned}',
15:r'\begin{aligned}E_g(V)&=\operatorname{Tr}[(\widehat\rho-\rho_0)H_\mu(V)]+E_R(V),\\H_\mu&\mapsto H_\mu+C(V)I:\qquad \Delta E_g=0.\end{aligned}',
17:r'\begin{aligned}c^2d\tau^2&=Ac^2dt^2-Bd\mathbf x^2,\\\frac{d\tau}{dt}&=\sqrt{A(1-\beta_{\mathrm{local}}^2)},\qquad k=\lfloor\tau/\Delta\tau_*\rfloor.\end{aligned}',
20:r'\begin{aligned}U_*&=e^{-iH\Delta\tau_*/\hbar},\qquad U_*v_j=e^{-i\theta_j}v_j,\\E_j&=\frac{\hbar}{\Delta\tau_*}(\theta_j+2\pi\ell_j),\qquad \ell_j\in\mathbb Z.\end{aligned}'
}
for i,new in replacements.items():
 old=eqs[i-1];en=en.replace('$$'+old+'$$','$$\n'+new+'\n$$');ko=ko.replace('$$'+old+'$$','$$\n'+new+'\n$$')
en=en.replace('Physical records are written as script R',r'Physical records are written as $\mathcal R$')
ko=ko.replace('물리 기록은 필기체 R로',r'물리 기록은 $\mathcal R$로')
for l,s in (('EN',en),('KO',ko)):names[l].write_text(s)

def fonts(style,face,size):
 style.font.name=face;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');rp.append(rf)
 for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),face)

for lang in ('EN','KO'):
 path=ROOT/f'WRRA_M_1_0_Downstream_Synthesis_{lang}.docx'
 subprocess.run(['pandoc',str(names[lang]),'-o',str(path)],check=True)
 d=Document(path);face='NanumGothic' if lang=='KO' else 'Liberation Sans'
 # A dedicated OpenType math font keeps tau, hbar and delimiters distinguishable.
 settings=d.settings.element
 mp=settings.find(qn('m:mathPr'))
 if mp is None:mp=OxmlElement('m:mathPr');settings.append(mp)
 mf=mp.find(qn('m:mathFont'))
 if mf is None:mf=OxmlElement('m:mathFont');mp.insert(0,mf)
 mf.set(qn('m:val'),'Latin Modern Math')
 for mr in d._element.xpath('.//m:r'):
  rp=mr.find(qn('w:rPr'))
  if rp is None:rp=OxmlElement('w:rPr');mr.insert(0,rp)
  rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
  for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),'Latin Modern Math')
 for sec in d.sections:
  sec.page_height=Cm(29.7);sec.page_width=Cm(21)
  sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(2.1);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.15)
  sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9)
 for sn in ('Normal','Body Text','First Paragraph'):
  if sn in d.styles:
   st=d.styles[sn];fonts(st,face,10.5);st.paragraph_format.line_spacing=1.16;st.paragraph_format.space_after=Pt(5 if lang=='EN' else 6)
   st.paragraph_format.widow_control=True
 for i in range(1,5):
  if f'Heading {i}' not in d.styles:d.styles.add_style(f'Heading {i}',WD_STYLE_TYPE.PARAGRAPH)
  st=d.styles[f'Heading {i}'];fonts(st,face,13 if i==1 else 12);st.font.bold=True
  st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
 fonts(d.styles['Title'],face,20);fonts(d.styles['Subtitle'],face,12)
 d.styles['Title'].paragraph_format.space_after=Pt(8);d.styles['Subtitle'].paragraph_format.space_after=Pt(8)
 for p in list(d.paragraphs):
  if p.style.name in ('Author','Date'):
   p._element.getparent().remove(p._element)
  elif p.style.name=='Heading 2':p.style=d.styles['Heading 1']
  elif p.style.name in ('Title','Subtitle'):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  if lang=='EN' and p.text=='References':p.paragraph_format.page_break_before=True
 # Native equations are numbered once in source order and kept intact.
 n=0
 for p in d.paragraphs:
  if p._element.xpath('./m:oMathPara'):
   n+=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
   if n in (6,11):
    # Short function arguments need ordinary glyph parentheses, not stretched
    # delimiter strokes that some PDF readers/OCR mistake for vertical bars.
    for de in list(p._element.xpath('.//m:d')):
     dp=de.find(qn('m:dPr'))
     bc=dp.find(qn('m:begChr'));ec=dp.find(qn('m:endChr'))
     if bc is None or ec is None or bc.get(qn('m:val'))!='(' or ec.get(qn('m:val'))!=')':continue
     parent=de.getparent();pos=parent.index(de)
     for ch in ('(',')'):
      rr=OxmlElement('m:r');rp2=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rp2.append(sty);rr.append(rp2)
      tt=OxmlElement('m:t');tt.text=ch;rr.append(tt)
      if ch=='(':parent.insert(pos,rr);pos+=1
      else:closing=rr
     for ee in de.findall(qn('m:e')):
      for child in list(ee):parent.insert(pos,child);pos+=1
     parent.insert(pos,closing);parent.remove(de)
   p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8)
   r=p.add_run(f'   ({n})');r.font.size=Pt(9.5)
   for mr in p._element.xpath('.//m:r'):
    rp=mr.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');mr.insert(0,rp)
    z=rp.find(qn('w:sz'))
    if z is None:z=OxmlElement('w:sz');rp.append(z)
    z.set(qn('w:val'),'21')
  if p.text.startswith(('Table ','표 ')):
   p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(5)
   for r in p.runs:r.bold=True;r.font.size=Pt(9.5)
  if re.match(r'^\[[1-8]\]',p.text):
   p.paragraph_format.space_after=Pt(4 if lang=='EN' else 5)
   for r in p.runs:r.font.size=Pt(9)
   # Long source addresses may wrap at word boundaries without overflowing.
   for r in p.runs:
    r.text=r.text.replace('/blob/','/\u200bblob/').replace('/b4aa','/\u200bb4aa').replace('/DOWNSTREAM','/\u200bDOWNSTREAM')
 widths=[[2.2,7.1,7.4],[1.5,6.8,8.4],[5.3,5.7,5.7],[3,7.2,6.5],[6.2,1.3,4.6,4.6],[5.2,3.1,8.4]]
 assert len(d.tables)==6
 for ti,t in enumerate(d.tables):
  t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  pr=t._tbl.tblPr
  borders=OxmlElement('w:tblBorders')
  for edge in ('top','left','bottom','right','insideH','insideV'):
   x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
  old=pr.find(qn('w:tblBorders'))
  if old is not None:pr.remove(old)
  pr.append(borders)
  for rowi,row in enumerate(t.rows):
   trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if rowi==0:trpr.append(OxmlElement('w:tblHeader'))
   for ci,c in enumerate(row.cells):
    c.width=Cm(widths[ti][ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
    for edge,val in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
     x=OxmlElement('w:'+edge);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');margin.append(x)
    tcpr.append(margin)
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2' if rowi==0 else 'FFFFFF');tcpr.append(sh)
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
     p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (ci==0 and ti==1) or (ti==4 and ci>0) else WD_ALIGN_PARAGRAPH.LEFT
     for r in p.runs:r.font.name=face;r.font.size=Pt(9);r.bold=rowi==0;r.font.color.rgb=RGBColor(0,0,0)
  for i,col in enumerate(t.columns):col.width=Cm(widths[ti][i])
 footer=d.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
 d.core_properties.author='Wonsik Choi; Jeongin Choi';d.core_properties.title=d.paragraphs[0].text
 d.core_properties.subject='WRRA M downstream final consolidated research edition 1.0'
 d.save(path);print(lang,'numbered display equations',n,'tables',len(d.tables))

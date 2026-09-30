from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "paper" / "WRRA_M_0_4_EN.docx"

NAVY = "18324A"
PALE_BLUE = "EAF2F8"
LIGHT_BORDER = "D9E1E8"
MID_GRAY = "5E6872"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table, color=LIGHT_BORDER, size="5", val="single") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), val)
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def fixed_table(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    style = tbl_pr.find(qn("w:tblStyle"))
    if style is not None:
        tbl_pr.remove(style)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for index, width in enumerate(widths):
        if index < len(grid_columns):
            grid_columns[index].set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is not None:
                tc_w.set(qn("w:w"), str(int(widths[index] * 1440)))
                tc_w.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(tag)


def set_run_font(run, size=9.0, bold=None, color=None, italic=None) -> None:
    run.font.name = "Aptos"
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), "Aptos")
    rfonts.set(qn("w:hAnsi"), "Aptos")
    rfonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def format_table(table, widths: list[float], header=True, font_size=8.3) -> None:
    fixed_table(table, widths)
    table_borders(table)
    if header:
        repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if header and row_index == 0:
                shade(cell, NAVY)
            elif row_index % 2 == 0:
                shade(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                p_style = paragraph._p.get_or_add_pPr().find(qn("w:pStyle"))
                if p_style is not None:
                    paragraph._p.pPr.remove(p_style)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.02
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=True if header and row_index == 0 else None,
                        color="FFFFFF" if header and row_index == 0 else None,
                    )


def prepend_text(paragraph, text: str) -> None:
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    rpr.append(fonts)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    paragraph._p.insert(0, run)


def flatten_hyperlinks(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for hyperlink in list(paragraph._p.findall(qn("w:hyperlink"))):
            index = paragraph._p.index(hyperlink)
            for child in list(hyperlink):
                hyperlink.remove(child)
                paragraph._p.insert(index, child)
                index += 1
            paragraph._p.remove(hyperlink)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4.5)

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(6)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(11.5)
    subtitle.font.italic = True
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 14.5, 14, 6),
        ("Heading 2", 11.7, 10, 4),
        ("Heading 3", 10.5, 8, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Body Text" in doc.styles:
        body = doc.styles["Body Text"]
        body.font.name = "Aptos"
        body.font.size = Pt(9.5)
        body.paragraph_format.line_spacing = 1.08
        body.paragraph_format.space_after = Pt(4.5)


def delete_first_row(table) -> None:
    table._tbl.remove(table.rows[0]._tr)


def main() -> None:
    doc = Document(DOCX)
    configure_styles(doc)
    flatten_hyperlinks(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.58)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)

    # Compact title metadata: remove the redundant Markdown header row.
    metadata = doc.tables[0]
    delete_first_row(metadata)
    fixed_table(metadata, [1.15, 5.75])
    table_borders(metadata, color="FFFFFF", size="0", val="nil")
    for row in metadata.rows:
        for index, cell in enumerate(row.cells):
            cell_margins(cell, top=35, bottom=35)
            for paragraph in cell.paragraphs:
                p_style = paragraph._p.get_or_add_pPr().find(qn("w:pStyle"))
                if p_style is not None:
                    paragraph._p.pPr.remove(p_style)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=8.6, bold=index == 0, color=MID_GRAY if index == 0 else None)

    format_table(doc.tables[1], [1.55, 5.35], header=True, font_size=8.1)
    format_table(doc.tables[2], [4.75, 2.15], header=True, font_size=8.5)
    format_table(doc.tables[3], [4.75, 2.15], header=True, font_size=8.5)
    format_table(doc.tables[4], [0.48, 1.42, 5.00], header=True, font_size=7.9)

    bullet_starts = {
        "The exact distinction between",
        "A conditional extension of",
        "A positive common-carrier",
        "One non-circular scalar",
        "Exact identities reducing",
        "An exact finite witness",
        "Explicit failure conditions",
        "Numerical response signatures",
        "Proof that the four-filter",
        "Three generations, Yukawa",
        "Gravity, dark matter",
        "the transported band",
        "carrier dependence must",
        "channel differences may",
        "full abstract rank",
        "The sixteenth channel is",
        "Extending the carrier splits",
        "The response Gram operator",
        "The response signature or",
        "An independently fixed carrier",
        "The metric, band, channel labels",
        "A broader admissible filter",
        "the squared-distance expansions",
        "the finite witness yields",
        "the four witness scores",
        "adding a zero-charge singlet",
        "an identity sixteen-channel",
    }
    numbered_starts = (
        "Fix one admissible sixteen-channel",
        "Verify one principal causal cone",
        "Fix the response extraction map",
        "Calculate all eight",
        "Construct the eight",
        "Publish ",
        "Apply the unchanged",
    )

    protocol_number = 1
    reference_number = 1
    in_references = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "Appendix B References":
            in_references = True
        if any(text.startswith(start) for start in bullet_starts):
            prepend_text(paragraph, "• ")
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.14)
            paragraph.paragraph_format.space_after = Pt(2.2)
        elif protocol_number <= 7 and text.startswith(numbered_starts[protocol_number - 1]):
            prepend_text(paragraph, f"{protocol_number}. ")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(2.5)
            protocol_number += 1
        elif in_references and text.startswith("Choi, W."):
            prepend_text(paragraph, f"[{reference_number}] ")
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            for run in paragraph.runs:
                set_run_font(run, size=8.2)
            reference_number += 1

        # Preserve equations together and keep ordinary paragraphs visually compact.
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph._p.xpath(".//m:oMath | .//m:oMathPara"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(2.5)
            paragraph.paragraph_format.space_after = Pt(4)

    # Author/date and copyright are quiet metadata, not headings.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() in {"Wonsik Choi", "1 October 2026"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_run_font(run, size=9.0, color=MID_GRAY)
        if paragraph.text.strip().startswith("Copyright ©"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            for run in paragraph.runs:
                set_run_font(run, size=8.2, color=MID_GRAY)

    # Keep the paired gap identities together as one theorem display.
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == "Then the 0.3 gaps satisfy the exact identities":
            paragraph.paragraph_format.keep_with_next = True
            if index + 1 < len(paragraphs):
                paragraphs[index + 1].paragraph_format.keep_with_next = True
            break

    props = doc.core_properties
    props.title = "WRRA-M 0.4 Common-Carrier Metric Origin of Filter Compatibility"
    props.subject = "Conditional fifteen-to-sixteen-channel common-carrier extension and exact selection-gap identities"
    props.author = "Wonsik Choi"
    props.keywords = "WRRA-M, common carrier, filter selection, response metric, right-handed neutrino"
    props.comments = "Version 0.4, 1 October 2026"

    doc.save(DOCX)


if __name__ == "__main__":
    main()

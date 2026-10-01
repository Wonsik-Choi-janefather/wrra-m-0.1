from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import format_wrra_m_0_4 as base


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "paper" / "WRRA_M_0_4_KO.docx"
FONT = "Noto Sans KR"


def set_style_font(style, size: float, bold: bool | None = None) -> None:
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold


def paragraph_style(doc: Document, name: str):
    return next(
        style
        for style in doc.styles
        if style.type == WD_STYLE_TYPE.PARAGRAPH and style.name == name
    )


def set_korean_run_font(run, size: float | None = None) -> None:
    if run._element.xpath(".//m:oMath | .//m:oMathPara"):
        return
    run.font.name = FONT
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)


def configure_korean_styles(doc: Document) -> None:
    normal = paragraph_style(doc, "Normal")
    set_style_font(normal, 9.6)
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.space_after = Pt(4.8)

    title = paragraph_style(doc, "Title")
    set_style_font(title, 21.0, True)
    title.paragraph_format.space_after = Pt(7)

    subtitle = paragraph_style(doc, "Subtitle")
    set_style_font(subtitle, 10.8)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor.from_string(base.MID_GRAY)
    subtitle.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 14.2, 14, 6),
        ("Heading 2", 11.6, 10, 4),
        ("Heading 3", 10.5, 8, 3),
    ):
        style = paragraph_style(doc, name)
        set_style_font(style, size, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    body_styles = [style for style in doc.styles if style.name == "Body Text"]
    if body_styles:
        body = body_styles[0]
        set_style_font(body, 9.6)
        body.paragraph_format.line_spacing = 1.14
        body.paragraph_format.space_after = Pt(4.8)


def main() -> None:
    doc = Document(DOCX)
    base.flatten_hyperlinks(doc)
    configure_korean_styles(doc)

    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.58)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)
        section.header_distance = Inches(0.25)
        section.footer_distance = Inches(0.25)

    metadata = doc.tables[0]
    base.delete_first_row(metadata)
    base.fixed_table(metadata, [1.15, 5.75])
    base.table_borders(metadata, color="FFFFFF", size="0", val="nil")
    for row in metadata.rows:
        for index, cell in enumerate(row.cells):
            base.cell_margins(cell, top=35, bottom=35)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_korean_run_font(run, 8.6)
                    run.bold = index == 0
                    if index == 0:
                        run.font.color.rgb = RGBColor.from_string(base.MID_GRAY)

    base.format_table(doc.tables[1], [1.62, 5.28], header=True, font_size=8.0)
    base.format_table(doc.tables[2], [4.75, 2.15], header=True, font_size=8.4)
    base.format_table(doc.tables[3], [4.75, 2.15], header=True, font_size=8.4)
    base.format_table(doc.tables[4], [0.48, 1.50, 4.92], header=True, font_size=7.7)

    reference_mode = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        # Display math has no ordinary paragraph text. Keep it centered, while
        # paragraphs containing both Korean prose and inline math remain left aligned.
        has_math = bool(paragraph._p.xpath(".//m:oMath | .//m:oMathPara"))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if has_math and not text else WD_ALIGN_PARAGRAPH.LEFT
        if has_math and not text:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(2.5)
            paragraph.paragraph_format.space_after = Pt(4)
        if not text:
            continue
        if text == "부록 B 참고문헌":
            reference_mode = True
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph._p.xpath(".//m:oMathPara") and not text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(2.5)
            paragraph.paragraph_format.space_after = Pt(4)
        if reference_mode and text.startswith("Choi, W."):
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
        if text in {"최원식 Wonsik Choi", "2026년 10월 1일"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_korean_run_font(run, 9.0)
                run.font.color.rgb = RGBColor.from_string(base.MID_GRAY)
        if text.startswith("Copyright ©"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            for run in paragraph.runs:
                set_korean_run_font(run, 8.2)
                run.font.color.rgb = RGBColor.from_string(base.MID_GRAY)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_korean_run_font(run)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_korean_run_font(run)

    # Keep short witness and claim tables intact across page boundaries.
    for table in (doc.tables[2], doc.tables[3], doc.tables[4]):
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1

    # Ordinary arrows in this prose chain must stay text, not Word math objects.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("0.2 선택된 필터의 산출"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                set_korean_run_font(run)

    props = doc.core_properties
    props.title = "WRRA-M 0.4 필터 적합성의 공통전달자 계량 기원"
    props.subject = "15채널에서 16채널로의 조건부 공통전달자 확장과 정확한 선택 격차 항등식"
    props.author = "최원식 Wonsik Choi"
    props.keywords = "WRRA-M, 공통전달자, 필터 선택, 응답 계량, 오른손잡이 중성미자"
    props.comments = "WRRA-M 0.4 한글판 2026년 10월 1일"

    doc.save(DOCX)


if __name__ == "__main__":
    main()

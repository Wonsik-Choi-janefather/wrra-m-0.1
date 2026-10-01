"""Build matching native-equation Word editions; PDF rendering is a separate step."""
from pathlib import Path
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
WIDTHS=[[2.14,1.1,1.85,1.85],[.65,1.1,1.25,1.3,1.3,1.34],
        [5.,1.94],[2.2,2.2,2.54],[2.5,1.6,2.84]]

def font(style,name,size):
    style.font.name=name
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR')

def build(lang):
    source=ROOT/f'source_{lang.lower()}.md'
    out=OUT/f'WRRA_M_0_6_{lang}.docx'
    subprocess.run(['pandoc',str(source),'-o',str(out),'--resource-path',str(ROOT)],check=True,cwd=ROOT)
    doc=Document(out);section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.72)
    section.left_margin=section.right_margin=Inches(.78)
    for name in ['Normal','Body Text','First Paragraph','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Table']:
        if name not in doc.styles:
            style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            style.base_style=doc.styles['Normal']
        st=doc.styles[name]
        font(st,'Liberation Serif',11)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.15
    for name,size in [('Title',20),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
        font(doc.styles[name],'Liberation Sans',size)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(13)
    tags=re.findall(r'\\tag\{(\d+)\}',source.read_text())
    equations=[p for p in doc.paragraphs if p._p.xpath('.//m:oMath')]
    assert tags==list(map(str,range(1,23))) and len(equations)==22
    for p,number in zip(equations,tags):
        run=OxmlElement('m:r');prop=OxmlElement('m:rPr');st=OxmlElement('m:sty')
        st.set(qn('m:val'),'p');prop.append(st);run.append(prop)
        t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text='   ('+number+')'
        run.append(t);p._p.xpath('.//m:oMath')[-1].append(run)
        p.paragraph_format.space_before=Pt(4)
        p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.keep_together=True
    in_refs=False
    for p in doc.paragraphs:
        p.paragraph_format.first_line_indent=Inches(0)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if p.style.name=='Heading 1' and p.text.startswith('WRRA'):p.style=doc.styles['Title']
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
        if p.style.name=='Caption':
            p.paragraph_format.space_after=Pt(10)
            font(doc.styles['Caption'],'Liberation Serif',10)
        if p.text in ['참고문헌과 재현 자료','References and reproduction material']:
            in_refs=True;continue
        if p.text.startswith(('compute.py','The calibrations')):in_refs=False
        if in_refs:
            p.paragraph_format.line_spacing=1.
            p.paragraph_format.space_after=Pt(2)
            for run in p.runs:run.font.size=Pt(10)
    paragraphs=doc.paragraphs
    for p,nxt in zip(paragraphs,paragraphs[1:]):
        if p.text and not p._p.xpath('.//m:oMath') and nxt._p.xpath('.//m:oMath'):
            p.paragraph_format.keep_with_next=True
    assert len(doc.tables)==len(WIDTHS)
    for table,cols in zip(doc.tables,WIDTHS):
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        for column,width in zip(table.columns,cols):
            column.width=Inches(width)
            for cell in column.cells:cell.width=Inches(width)
        for col,width in zip(table._tbl.tblGrid.gridCol_lst,cols):col.set(qn('w:w'),str(round(width*1440)))
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+side)
            for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+key),value)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for ri,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcp=cell._tc.get_or_add_tcPr();shading=OxmlElement('w:shd')
                shading.set(qn('w:fill'),'E5EBF0' if ri==0 else ('F7F8FA' if ri%2==0 else 'FFFFFF'));tcp.append(shading)
                padding=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');padding.append(el)
                tcp.append(padding)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.1
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(cols) in [2,3] or ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name='Liberation Serif';run.font.size=Pt(10);run.bold=ri==0
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='Wonsik Choi'
    doc.core_properties.title='WRRA-M 0.6 Information Load and Twist Gravity with Expansion in a Finite Model'
    doc.core_properties.subject=lang+' edition with matched equations and calculations'
    doc.save(out);print(out)

if __name__=='__main__':
    build('KO');build('EN')

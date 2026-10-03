"""Build matched native-equation Word papers using the established 0.7 design."""
from pathlib import Path
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent.parent/'paper'


def font(style,name,size):
    style.font.name=name
    style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'NanumGothic')


def build(lang):
    source=ROOT/f'source_{lang.lower()}.md'
    out=OUT/f'WRRA_M_0_12_{lang}.docx'
    OUT.mkdir(exist_ok=True)
    subprocess.run(['pandoc',str(source),'-o',str(out)],check=True)
    doc=Document(out)
    section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.72)
    section.left_margin=section.right_margin=Inches(.78)
    for name in ['Normal','Body Text','First Paragraph','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','Table']:
        if name not in doc.styles:
            st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            st.base_style=doc.styles['Normal']
        st=doc.styles[name]
        font(st,'Liberation Serif',11)
        st.paragraph_format.space_after=Pt(7)
        st.paragraph_format.line_spacing=1.13
        if lang=='KO' and name in ('Normal','Body Text','First Paragraph'):
            st.paragraph_format.line_spacing=1.08
            st.paragraph_format.space_after=Pt(5)
        if lang=='EN' and name in ('Normal','Body Text','First Paragraph'):
            st.paragraph_format.line_spacing=1.05
            st.paragraph_format.space_after=Pt(4)
    for name,size in [('Title',20),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
        font(doc.styles[name],'Liberation Sans',size)
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(13)
    equations=[p for p in doc.paragraphs if p._p.xpath('.//m:oMath')]
    tags=re.findall(r'\\tag\{(\d+)\}',source.read_text())
    assert tags==list(map(str,range(1,19))) and len(equations)==18
    for p,number in zip(equations,tags):
        run=OxmlElement('m:r'); prop=OxmlElement('m:rPr'); st=OxmlElement('m:sty')
        st.set(qn('m:val'),'p'); prop.append(st); run.append(prop)
        t=OxmlElement('m:t'); t.set(qn('xml:space'),'preserve'); t.text='   ('+number+')'
        run.append(t); p._p.xpath('.//m:oMath')[-1].append(run)
        p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.keep_together=True
    for p in doc.paragraphs:
        p.paragraph_format.first_line_indent=Inches(0)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if p.style.name=='Heading 1' and p.text.startswith('WRRA'):
            p.style=doc.styles['Title']
        if p.style.name.startswith('Heading') or p.style.name=='Title':
            for run in p.runs: run.font.color.rgb=RGBColor(0,0,0)
    for p,nxt in zip(doc.paragraphs,doc.paragraphs[1:]):
        if p.text and not p._p.xpath('.//m:oMath') and nxt._p.xpath('.//m:oMath'):
            p.paragraph_format.keep_with_next=True
    in_refs=False
    for p in doc.paragraphs:
        if p.text.startswith(('Choi Wonsik.','WRRA-M repository.')):
            in_refs=True
        if in_refs:
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.space_after=Pt(3)
            for run in p.runs:run.font.size=Pt(10)
    # Keep the reference heading and its short bibliography together.
    in_refs=False
    for p in doc.paragraphs:
        if p.text in ('References','참고 자료'): in_refs=True
        if in_refs and not p.text.startswith('Copyright'):
            p.paragraph_format.keep_with_next=True
    for p in reversed(doc.paragraphs):
        if p.text.startswith('Hughes Scott.'):
            p.paragraph_format.keep_with_next=False
            break
    widths=[[1.5,3.3,2.14],[.7,1.85,1.85,2.54],[.8,3.05,3.09],[1.5,.5,2.47,2.47],[2.35,3.39,1.2]]
    assert len(doc.tables)==5
    for ti,(table,cols) in enumerate(zip(doc.tables,widths)):
        compact=ti==2 or (ti==0 and lang=='EN')
        table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        for column,width in zip(table.columns,cols):
            column.width=Inches(width)
            for cell in column.cells: cell.width=Inches(width)
        for col,width in zip(table._tbl.tblGrid.gridCol_lst,cols):
            col.set(qn('w:w'),str(round(width*1440)))
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+side)
            for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+key),value)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for ri,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if ri==0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcp=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
                sh.set(qn('w:fill'),'E5EBF0' if ri==0 else ('F7F8FA' if ri%2==0 else 'FFFFFF'));tcp.append(sh)
                padding=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side); el.set(qn('w:w'),'45' if compact and side in ('top','bottom') else '90'); el.set(qn('w:type'),'dxa'); padding.append(el)
                tcp.append(padding)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(1 if compact else 3); p.paragraph_format.space_after=Pt(1 if compact else 3)
                    p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name='Liberation Serif'; run.font.size=Pt(10); run.bold=ri==0
    # Keep the license with every page instead of spilling it onto an empty last page.
    license_text='Copyright 2026 Wonsik Choi. CC BY 4.0. https://creativecommons.org/licenses/by/4.0/'
    for p in list(doc.paragraphs):
        if p.text.startswith('Copyright 2026 Wonsik Choi.'):
            p._element.getparent().remove(p._element)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_after=Pt(0)
    license_run=footer.add_run(license_text+'     ')
    license_run.font.name='Liberation Serif';license_run.font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='Wonsik Choi'
    doc.core_properties.title='WRRA-M 0.12 Proper-Time Updates and Mode Spectra under Finite Boundaries'
    doc.core_properties.subject=lang+' edition with matched equations and calculated accounting checks'
    doc.save(out)
    (OUT/f'WRRA_M_0_12_{lang}.md').write_text(source.read_text())
    print(out)


if __name__=='__main__':
    build('KO');build('EN')

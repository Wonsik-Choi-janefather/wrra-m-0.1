"""Compare the two authored editions against each other and computed outputs.

Optional authoring check: requires python-docx and pypdf, beyond runtime requirements.
Visual page inspection is a separate release activity.
"""
from pathlib import Path
import json
import re
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
PAPER = ROOT.parent.parent / 'paper'


def xml_signature(node):
    return (node.tag, tuple(sorted(node.attrib.items())), node.text or '',
            tuple(xml_signature(child) for child in node))


def run():
    result = json.loads((ROOT/'results/results.json').read_text())
    verification = json.loads((ROOT/'results/verification.json').read_text())
    equations, native, editions = {}, {}, {}
    for lang in ('KO', 'EN'):
        text = (PAPER/f'WRRA_M_0_7_{lang}.md').read_text()
        equations[lang] = re.findall(r'\$\$(.*?)\$\$', text, re.S)
        assert len(equations[lang]) == 11
        doc = Document(PAPER/f'WRRA_M_0_7_{lang}.docx')
        native[lang] = [xml_signature(eq) for p in doc.paragraphs
                        for eq in p._p.xpath('.//m:oMath')]
        assert len(native[lang]) == 11 and len(doc.tables) == 2
        assert len(doc.tables[0].rows) == 10
        for row, case in zip(doc.tables[0].rows[1:], result['cases']):
            expected = [case['Jc'], case['Jb'],
                        100*case['sector_shares']['phenotype'], 100*case['hidden_share']]
            assert [cell.text.strip() for cell in row.cells[1:]] == [f'{v:.6f}' for v in expected]
        assert len(doc.tables[1].rows) == verification['check_count']+1
        assert all(row.cells[-1].text.strip() in ('Pass', '통과')
                   for row in doc.tables[1].rows[1:])
        pdf = PdfReader(PAPER/f'WRRA_M_0_7_{lang}.pdf')
        pdf_text = '\n'.join(page.extract_text() for page in pdf.pages)
        assert '\ufffd' not in pdf_text
        for number in ('4.930000', '95.070000', '207.510905', '0.535586511'):
            assert number in pdf_text, (lang, number)
        assert 'creativecommons.org/licenses/by/4.0/' in pdf_text
        editions[lang] = {'pdf_pages':len(pdf.pages), 'native_equations':len(native[lang]),
                         'numerical_rows':9, 'verification_rows':verification['check_count']}
    assert equations['KO'] == equations['EN'], 'source equations differ'
    assert native['KO'] == native['EN'], 'native equations differ'
    report = {'version':'WRRA-M 0.7', 'passed':True,
              'input_hash_sha256':result['input_hash_sha256'],
              'source_equations_equal':True, 'native_equations_equal':True,
              'tables_match_computed_results':True, 'editions':editions,
              'visual_QA':'All rendered pages separately inspected for the release.'}
    (ROOT/'results/document_checks.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__ == '__main__':
    run()

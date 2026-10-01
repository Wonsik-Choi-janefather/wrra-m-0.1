from pathlib import Path
import subprocess
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
source = ROOT / "source_en.md"
text = source.read_text()
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)
docx = OUT / "WRRA_M_0_5_EN.docx"
subprocess.run(
    ["pandoc", str(source), "-o", str(docx), "--resource-path", str(ROOT)],
    check=True, cwd=ROOT
)
doc = Document(docx)
for name in ["Normal", "Body Text", "First Paragraph", "Title", "Subtitle",
             "Heading 1", "Heading 2", "Heading 3", "Caption"]:
    if name not in doc.styles:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.72)
section.left_margin = section.right_margin = Inches(.78)
for name in ["Normal", "Body Text", "First Paragraph", "Title", "Subtitle",
             "Heading 1", "Heading 2", "Heading 3", "Caption", "Table"]:
    if name not in doc.styles:
        continue
    style = doc.styles[name]
    style.font.name = "Liberation Serif"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.15
for name, size in [("Title", 20), ("Heading 1", 17),
                   ("Heading 2", 14), ("Heading 3", 12)]:
    doc.styles[name].font.name = "Liberation Sans"
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles["Heading 2"].paragraph_format.space_before = Pt(13)
for p in doc.paragraphs:
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if p.style.name == "Heading 1" and p.text.startswith("WRRA"):
        p.style = doc.styles["Title"]
    if p.text.rstrip().endswith(":") or p.text.rstrip().endswith("then"):
        p.paragraph_format.keep_with_next = True
    if p.text.startswith("Rotation speeds of the spherical"):
        p.style = doc.styles["Caption"]
        p.paragraph_format.space_after = Pt(10)

tags = re.findall(r"\\tag\{(\d+)\}", text)
equations = [p for p in doc.paragraphs if p._p.xpath(".//m:oMath")]
assert tags == [str(i) for i in range(1, 16)]
assert len(equations) == 15
for p, number in zip(equations, tags):
    math = p._p.xpath(".//m:oMath")[-1]
    run = OxmlElement("m:r")
    prop = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    prop.append(style)
    run.append(prop)
    t = OxmlElement("m:t")
    t.set(qn("xml:space"), "preserve")
    t.text = "   (" + number + ")"
    run.append(t)
    math.append(run)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)

widths = [
    [1.30, 5.64],
    [.75, 1.20, 1.20, 1.69, 2.10],
    [1.30, 1.64, 2.00, 2.00],
    [1.00, 1.70, 1.70, 2.54],
    [1.22, 1.10, 1.46, 1.64, 1.52],
    [2.42, 2.10, 2.42],
]
assert len(doc.tables) == len(widths)
for table, cols in zip(doc.tables, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, cols):
        column.width = Inches(width)
        for cell in column.cells:
            cell.width = Inches(width)
    for col, width in zip(table._tbl.tblGrid.gridCol_lst, cols):
        col.set(qn("w:w"), str(round(width * 1440)))
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for ri, row in enumerate(table.rows):
        rowpr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        rowpr.append(cant_split)
        if ri == 0:
            rowpr.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E5EBF0" if ri == 0 else
                        ("F7F8FA" if ri % 2 == 0 else "FFFFFF"))
            tcp.append(shading)
            padding = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), "90")
                item.set(qn("w:type"), "dxa")
                padding.append(item)
            tcp.append(padding)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.1
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (
                    len(cols) == 2 or len(cols) == 3 or ci == 0
                ) else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Liberation Serif"
                    run.font.size = Pt(10)
                    run.bold = ri == 0
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
in_references = False
for p in doc.paragraphs:
    if p.text == "References and reproduction material":
        in_references = True
        continue
    if p.text.startswith("In the reproduction package"):
        in_references = False
    if in_references:
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(10)
doc.core_properties.author = "Wonsik Choi"
doc.core_properties.title = (
    "WRRA-M 0.5 Information Load and Twist Gravity in a Finite Model"
)
doc.core_properties.subject = (
    "English edition of the Korean WRRA-M 0.5 paper with identical calculations"
)
doc.save(docx)
print(docx)
